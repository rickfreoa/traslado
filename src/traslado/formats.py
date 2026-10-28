# SPDX-License-Identifier: AGPL-3.0-or-later
# Copyright (C) 2026 Ricardo Peres Freoa
from __future__ import annotations

import csv
import html
from html.parser import HTMLParser
import io
import json
from pathlib import Path
import re
import warnings

from .errors import ConversionError, exception_name, localized

from .model import Block, Document, Options, Page, text_metrics
from .ocr import recognize
from .safety import check_archive

from .registry import SUPPORTED, lookup


class HTMLText(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style", "noscript"}:
            self.skip += 1
        elif not self.skip and tag in {"p", "div", "br", "tr", "h1", "h2", "h3", "li"}:
            self.parts.append("\n")
        elif not self.skip and tag in {"td", "th"}:
            self.parts.append("\t")

    def handle_endtag(self, tag):
        if tag in {"script", "style", "noscript"} and self.skip:
            self.skip -= 1
        elif not self.skip and tag in {"p", "div", "tr", "h1", "h2", "h3", "li"}:
            self.parts.append("\n")

    def handle_data(self, data):
        if not self.skip:
            self.parts.append(data)


def image_to_png(data: bytes | Path, output: Path, options: Options) -> None:
    from PIL import Image
    source = io.BytesIO(data) if isinstance(data, bytes) else data
    with warnings.catch_warnings():
        warnings.simplefilter("error", Image.DecompressionBombWarning)
        with Image.open(source) as im:
            if im.width * im.height > options.max_pixels:
                raise ConversionError('resource_limit', "Imagem excede limite de pixels")
            _load_image(im)
            im.convert("RGB").save(output, "PNG")


def _load_image(image) -> None:
    try:
        image.load()
    except OSError as exc:
        # Pillow reports truncated/invalid raster data as OSError without errno.
        # Real filesystem failures retain their errno and their I/O code.
        if exc.errno is None:
            error = ConversionError("malformed", str(exc))
            error._warning_name = type(exc).__name__
            raise error from exc
        raise


def _docx_paragraph_text(paragraph, note_reference=None) -> str:
    from docx.oxml.ns import qn
    parts = []
    pending = []

    def walk(node):
        for child in node:
            if child.tag == qn("w:r"):
                # The custom mark belongs to the following run. Resolve now
                # (for first-call order), but write after that run's text.
                previous = pending[:]
                pending.clear()
                walk(child)
                parts.extend(previous)
            elif child.tag in {qn("w:footnoteReference"), qn("w:endnoteReference")} and note_reference:
                family = "footnote" if child.tag == qn("w:footnoteReference") else "endnote"
                marker = note_reference(family, child.get(qn("w:id"), ""))
                if any(n.get(qn("w:customMarkFollows")) in {"1", "true", "on"}
                       for n in (child, node)):
                    pending.append(marker)
                else:
                    parts.append(marker)
            elif child.tag == qn("w:t"):
                parts.append(child.text or "")
            elif child.tag == qn("w:tab"):
                parts.append("\t")
            elif child.tag in {qn("w:br"), qn("w:cr")}:
                parts.append("\n")
            elif child.tag != qn("w:p"):
                # A descendant paragraph is visited separately by the cell walker.
                walk(child)

    walk(paragraph)
    parts.extend(pending)
    return "".join(parts)


def _docx_body_blocks(parent):
    from docx.oxml.ns import qn
    for node in parent.iterchildren():
        if node.tag in {qn("w:p"), qn("w:tbl")}:
            yield node
        elif node.tag == qn("w:sdt"):
            content = node.find(qn("w:sdtContent"))
            if content is not None:
                yield from _docx_body_blocks(content)


def docx_table(table, options: Options, note_reference=None) -> Block:
    from docx.oxml.ns import qn
    rows = []
    count = 0
    notes = []
    for row in table._tbl.tr_lst:
        cells = []
        for tc in row.tc_lst:
            count += 1
            if count > options.max_cells:
                raise ConversionError('resource_limit', "Tabela excede limite de células")
            pr = tc.tcPr
            span = 1
            merge = None
            if pr is not None:
                grid = pr.find(qn("w:gridSpan"))
                vertical = pr.find(qn("w:vMerge"))
                span = int(grid.get(qn("w:val"), "1")) if grid is not None else 1
                merge = vertical.get(qn("w:val"), "continue") if vertical is not None else None
            if not 1 <= span <= options.max_cells:
                raise ConversionError('resource_limit', "Mescla de tabela excede limite de células")
            text = "\n".join(_docx_paragraph_text(p, note_reference) for p in tc.iter(qn("w:p")))
            if any(True for _ in tc.iter(qn("w:tbl"))):
                notes.append("docx_nested_table_structure_flattened")
            cells.append({"text": text, "span": span, "merge": merge, "rowspan": 1, "skip": False})
        before = row.trPr.find(qn("w:gridBefore")) if row.trPr is not None else None
        start = int(before.get(qn("w:val"), "0")) if before is not None else 0
        if not 0 <= start <= options.max_cells:
            raise ConversionError('malformed', "Deslocamento de tabela inválido")
        rows.append((start, cells))
    active = {}
    for start, row in rows:
        col = start
        next_active = {}
        for cell in row:
            key = (col, cell["span"])
            if cell["merge"] == "continue" and key in active:
                anchor = active[key]
                anchor["rowspan"] += 1
                if cell["text"].strip():
                    anchor["text"] += "\n" + cell["text"]
                    notes.append("docx_vmerge_continuation_text_preserved")
                cell["skip"] = True
                next_active[key] = anchor
            elif cell["merge"] == "restart":
                next_active[key] = cell
            elif cell["merge"] is not None:
                # Orphans and incompatible spans remain independent visible cells.
                notes.append("docx_vmerge_inconsistent_content_preserved")
            col += cell["span"]
        # Only merges present in the immediately preceding row can continue.
        active = next_active
    rendered = ["<table>"]
    for start, row in rows:
        rendered.append("<tr>")
        if start:
            rendered.append(f'<td colspan="{start}"></td>')
        for cell in row:
            if not cell["skip"]:
                rendered.append(f'<td colspan="{cell["span"]}" rowspan="{cell["rowspan"]}">' +
                                html.escape(cell["text"]).replace("\n", "<br>") + "</td>")
        rendered.append("</tr>")
    rendered.append("</table>")
    return Block("html_table", "\n".join(rendered), engine="python-docx",
                 details={"warnings": sorted(set(notes))})


def _docx_note_text(note) -> tuple[str, bool]:
    """Flatten note content in XML order, including paragraphs in text boxes."""
    from docx.oxml.ns import qn
    paragraphs = []
    parts = []
    complex_content = False

    def flush():
        if parts:
            paragraphs.append("".join(parts))
            parts.clear()

    def walk(node):
        nonlocal complex_content
        if node.tag in {qn("w:tbl"), qn("w:drawing"), qn("w:pict"), qn("w:object"), qn("w:txbxContent")}:
            complex_content = True
        if node.tag == qn("w:r") and any(c.tag in {qn("w:footnoteRef"), qn("w:endnoteRef")} for c in node):
            return
        if node.tag == qn("w:p"):
            flush()
            paragraph_start = len(paragraphs)
        if node.tag == qn("w:t"):
            parts.append(node.text or "")
        elif node.tag == qn("w:tab"):
            parts.append("\t")
        elif node.tag in {qn("w:br"), qn("w:cr")}:
            parts.append("\n")
        else:
            for child in node:
                walk(child)
        if node.tag == qn("w:p"):
            if parts:
                flush()
            elif len(paragraphs) == paragraph_start:
                paragraphs.append("")

    walk(note)
    flush()
    return "\n\n    ".join(paragraphs), complex_content


def _docx_related_notes(doc, page):
    from docx.oxml import parse_xml
    from docx.oxml.ns import qn
    definitions = {"footnote": {}, "endnote": {}, "_extra_footnote": [], "_extra_endnote": []}
    for relation in doc.part.rels.values():
        if relation.is_external:
            continue
        kind = relation.reltype.rsplit("/", 1)[-1]
        if kind not in {"footnotes", "endnotes", "header", "footer", "comments"}:
            continue
        root = parse_xml(relation.target_part.blob)
        if kind in {"footnotes", "endnotes"}:
            family = kind[:-1]
            for note in root:
                if note.tag != qn("w:" + family) or note.get(qn("w:type")) in {
                        "separator", "continuationSeparator", "continuationNotice"}:
                    continue
                text, complex_content = _docx_note_text(note)
                number = note.get(qn("w:id"))
                if not number or number in definitions[family]:
                    # No definition is dropped: a missing or repeated id cannot be called unambiguously,
                    # so the text is kept as an extra, uncalled definition with its own warning.
                    definitions["_extra_" + family].append((number or "", text))
                    if "docx_note_duplicate_or_missing_id" not in page.warnings:
                        page.warnings.append("docx_note_duplicate_or_missing_id")
                else:
                    definitions[family][number] = text
                if complex_content and "docx_note_complex_content_flattened" not in page.warnings:
                    page.warnings.append("docx_note_complex_content_flattened")
        elif kind in {"header", "footer"}:
            if any((t.text or "").strip() for t in root.iter(qn("w:t"))):
                if "docx_headers_footers_not_extracted" not in page.warnings:
                    page.warnings.append("docx_headers_footers_not_extracted")
        elif any(True for _ in root.iter(qn("w:comment"))):
            if "docx_comments_not_extracted" not in page.warnings:
                page.warnings.append("docx_comments_not_extracted")
    return definitions


def _extract_docx(path: Path, workspace: Path, options: Options, document: Document) -> None:
    suffix = path.suffix.lower()
    assets = workspace / "assets"
    page = Page(1, "native")
    document.pages.append(page)
    check_archive(path, options)
    from docx import Document as Docx
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    doc = Docx(path)
    definitions = _docx_related_notes(doc, page)
    called = {"footnote": {}, "endnote": {}}

    def note_reference(family, number):
        called[family][number] = None
        if number not in definitions[family]:
            if "docx_note_reference_unresolved" not in page.warnings:
                page.warnings.append("docx_note_reference_unresolved")
            label = "nota" if family == "footnote" else "nota de fim"
            return f"[{label} {number} não encontrada]"
        return _note_label(family, number)

    asset_index = 0
    if any(True for _ in doc.element.body.iter(qn("w:ins"))) or any(True for _ in doc.element.body.iter(qn("w:del"))):
        page.warnings.append("docx_tracked_changes_not_resolved")
    for node in _docx_body_blocks(doc.element.body):
        if node.tag == qn("w:p"):
            paragraph = Paragraph(node, doc)
            text = _docx_paragraph_text(node, note_reference)
            if text.strip():
                style = paragraph.style.name.lower() if paragraph.style else ""
                prefix = ""
                if style.startswith("heading ") and style[-1:].isdigit():
                    prefix = "#" * max(1, min(6, int(style[-1]))) + " "
                page.blocks.append(Block("text", prefix + text, engine="python-docx"))
            if options.images:
                for blip in node.iter(qn("a:blip")):
                    rid = blip.get(qn("r:embed"))
                    if rid and rid in doc.part.related_parts:
                        asset_index += 1
                        target = assets / f"docx-{asset_index}.png"
                        try:
                            image_to_png(doc.part.related_parts[rid].blob, target, options)
                            page.blocks.append(Block("image", asset=f"assets/{target.name}", engine="python-docx"))
                        except (OSError, ValueError) as exc:
                            localized(document, page, exc)
                            page.warnings.append("docx_image_failed: " + exception_name(exc))
        elif node.tag == qn("w:tbl"):
            block = docx_table(Table(node, doc), options, note_reference)
            page.blocks.append(block)
            page.warnings.extend(block.details.get("warnings", []))
    for family in ("footnote", "endnote"):
        order = list(called[family]) + [n for n in definitions[family] if n not in called[family]]
        for number in order:
            if number not in definitions[family]:
                continue
            if number not in called[family] and "docx_note_definition_unreferenced" not in page.warnings:
                page.warnings.append("docx_note_definition_unreferenced")
            page.blocks.append(Block("text", _note_label(family, number) + ": " + definitions[family][number],
                                     engine="python-docx", details={"note_family": family, "note_id": number}))
        for serial, (number, text) in enumerate(definitions["_extra_" + family], 1):
            label = _note_label(family, (number or "sem-id") + f"-extra{serial}")
            page.blocks.append(Block("text", label + ": " + text, engine="python-docx",
                                     details={"note_family": family, "note_id": number, "note_extra": serial}))
    page.warnings.append("docx_formatting_partial")


def _note_label(family: str, number: str) -> str:
    """Markdown footnote label; the raw Word id stays in details.note_id."""
    safe = re.sub(r"[^A-Za-z0-9_-]", "_", number) or "sem-id"
    return "[^" + ("e" if family == "endnote" else "") + safe + "]"


def _extract_xlsx(path: Path, workspace: Path, options: Options, document: Document) -> None:
    suffix = path.suffix.lower()
    assets = workspace / "assets"
    page = Page(1, "native")
    document.pages.append(page)
    check_archive(path, options)
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    document.pages.clear()
    total_cells = 0
    try:
        if len(wb.worksheets) > options.max_pages:
            raise ConversionError('resource_limit', "Planilha excede limite de abas")
        for index, sheet in enumerate(wb.worksheets, 1):
            rows = []
            for row in sheet.iter_rows(values_only=True):
                total_cells += len(row)
                if total_cells > options.max_cells:
                    raise ConversionError('resource_limit', "Planilha excede limite de células")
                rows.append(["" if c is None else str(c) for c in row])
            document.pages.append(Page(index, "native", [Block("text", f"## {sheet.title}"),
                                  Block("table", rows=rows, engine="openpyxl")],
                                  ["xlsx_formulas_preserved_not_calculated", "xlsx_merges_styles_images_not_extracted"]))
    finally:
        wb.close()


def _extract_pptx(path: Path, workspace: Path, options: Options, document: Document) -> None:
    suffix = path.suffix.lower()
    assets = workspace / "assets"
    page = Page(1, "native")
    document.pages.append(page)
    check_archive(path, options)
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    prs = Presentation(path)
    if len(prs.slides) > options.max_pages:
        raise ConversionError('resource_limit', "Apresentação excede limite de slides")
    document.pages.clear()
    for index, slide in enumerate(prs.slides, 1):
        page = Page(index, "native", warnings=["pptx_groups_notes_charts_not_extracted", "pptx_order_geometric_heuristic"])
        document.pages.append(page)
        shapes = list(slide.shapes)
        if any(shape.top is None or shape.left is None for shape in shapes):
            # Missing transforms have no trustworthy geometric position.
            # Preserve source shape order rather than invent coordinates.
            page.warnings.append("pptx_shape_coordinates_missing_document_order")
        else:
            shapes.sort(key=lambda shape: (shape.top, shape.left))
        for number, shape in enumerate(shapes, 1):
            if shape.has_table:
                rows = [[c.text for c in row.cells] for row in shape.table.rows]
                if sum(len(r) for r in rows) > options.max_cells:
                    raise ConversionError('resource_limit', "Tabela excede limite de células")
                page.blocks.append(Block("table", rows=rows, engine="python-pptx"))
            elif shape.has_text_frame:
                page.blocks.append(Block("text", shape.text, engine="python-pptx"))
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE and options.images:
                target = assets / f"slide-{index}-{number}.png"
                try:
                    image_to_png(shape.image.blob, target, options)
                except (OSError, ValueError) as exc:
                    localized(document, page, exc)
                    # A failed optional image cannot discard valid slide
                    # text/tables. Never publish an incomplete image asset.
                    target.unlink(missing_ok=True)
                    page.warnings.append(f"pptx_image_failed: shape={number} {exception_name(exc)}")
                else:
                    page.blocks.append(Block("image", asset=f"assets/{target.name}", engine="python-pptx"))


def _extract_image(path: Path, workspace: Path, options: Options, document: Document) -> None:
    suffix = path.suffix.lower()
    assets = workspace / "assets"
    page = Page(1, "native")
    document.pages.append(page)
    from PIL import Image
    document.pages.clear()
    with Image.open(path) as source:
        count = getattr(source, "n_frames", 1)
        if count > options.max_pages:
            raise ConversionError('resource_limit', "Imagem excede limite de quadros")
        for index in range(count):
            source.seek(index)
            if source.width * source.height > options.max_pixels:
                raise ConversionError('resource_limit', "Imagem excede limite de pixels")
            _load_image(source)
            target = assets / f"image-{index+1}.png"
            source.convert("RGB").save(target, "PNG")
            page = Page(index+1, "image" if options.mode == "native" else "ocr")
            document.pages.append(page)
            if options.images:
                page.blocks.append(Block("image", asset=f"assets/{target.name}", engine="pillow"))
            if options.mode == "native":
                page.warnings.append("image_text_not_extracted_native_mode")
            else:
                try:
                    block, notes = recognize(target, workspace, options)
                    page.blocks.append(block)
                    page.warnings.extend(notes)
                except (RuntimeError, TimeoutError, ValueError, OSError) as exc:
                    localized(document, page, exc)
                    page.warnings.append("ocr_failed: " + str(exc))
            page.warnings.append("image_semantics_unverified")
            if not options.images:
                target.unlink(missing_ok=True)


def _extract_csv(path: Path, workspace: Path, options: Options, document: Document) -> None:
    suffix = path.suffix.lower()
    assets = workspace / "assets"
    page = Page(1, "native")
    document.pages.append(page)
    try:
        text = path.read_text(encoding="utf-8-sig", errors="strict")
    except UnicodeDecodeError as exc:
        raise ConversionError('malformed', "Texto deve estar em UTF-8; converta a codificação antes de processar") from exc
    reader = csv.reader(io.StringIO(text, newline=""), delimiter="\t" if suffix == ".tsv" else ",")
    rows, count = [], 0
    for row in reader:
        count += len(row)
        if count > options.max_cells:
            raise ConversionError('resource_limit', "CSV excede limite de células")
        rows.append(row)
    page.blocks.append(Block("table", rows=rows, engine="csv"))
    metrics = text_metrics(text)
    page.metrics = metrics
    if metrics["suspect"]:
        page.warnings.append("unicode_suspect")


def _extract_html(path: Path, workspace: Path, options: Options, document: Document) -> None:
    suffix = path.suffix.lower()
    assets = workspace / "assets"
    page = Page(1, "native")
    document.pages.append(page)
    try:
        text = path.read_text(encoding="utf-8-sig", errors="strict")
    except UnicodeDecodeError as exc:
        raise ConversionError('malformed', "Texto deve estar em UTF-8; converta a codificação antes de processar") from exc
    parser = HTMLText()
    parser.feed(text)
    page.blocks.append(Block("text", "".join(parser.parts), engine="html.parser"))
    page.warnings.append("html_table_structure_not_preserved")
    metrics = text_metrics(text)
    page.metrics = metrics
    if metrics["suspect"]:
        page.warnings.append("unicode_suspect")


def _extract_text(path: Path, workspace: Path, options: Options, document: Document) -> None:
    suffix = path.suffix.lower()
    assets = workspace / "assets"
    page = Page(1, "native")
    document.pages.append(page)
    try:
        text = path.read_text(encoding="utf-8-sig", errors="strict")
    except UnicodeDecodeError as exc:
        raise ConversionError('malformed', "Texto deve estar em UTF-8; converta a codificação antes de processar") from exc
    if suffix in {".json", ".xml"}:
        fence = "`" * max(3, 1 + max((len(x) for x in __import__('re').findall(r'`+', text)), default=0))
        page.blocks.append(Block("text", f"{fence}{suffix[1:]}\n{text}\n{fence}"))
    else:
        page.blocks.append(Block("text", text))
    metrics = text_metrics(text)
    page.metrics = metrics
    if metrics["suspect"]:
        page.warnings.append("unicode_suspect")


def _extract_pdf(path: Path, workspace: Path, options: Options, document: Document) -> None:
    from .pdf import extract as extract_pdf
    extract_pdf(path, workspace, options, document)


def extract(path: Path, workspace: Path, options: Options, document: Document) -> None:
    converter = lookup(path.suffix.lower())
    assets = workspace / "assets"
    assets.mkdir(exist_ok=True)
    if converter is None:
        raise ConversionError('unsupported', "Formato não suportado; use PDF, DOCX, XLSX, PPTX, texto, CSV, HTML ou imagem")
    try:
        converter.extract(path, workspace, options, document)
    except ConversionError:
        raise
    except Exception as exc:
        # Parser-specific failures are classified here, while filesystem errors
        # retain their I/O meaning. Unknown programming errors stay internal.
        import zipfile
        import zlib
        from xml.etree.ElementTree import ParseError
        if isinstance(exc, ModuleNotFoundError) and (exc.name or "").split(".")[0] in converter.requires:
            raise ConversionError("unsupported", str(exc)) from exc
        if isinstance(exc, ModuleNotFoundError) and (exc.name or "").split(".")[0] in converter.optional:
            raise ConversionError("missing_part", str(exc)) from exc
        if isinstance(exc, (zipfile.BadZipFile, zlib.error, ParseError)) or type(exc).__name__ in {
                "XMLSyntaxError", "FileDataError", "EmptyFileError", "UnidentifiedImageError"}:
            raise ConversionError("malformed", str(exc)) from exc
        if converter.id in {"docx", "xlsx", "pptx"}:
            required = {"docx": "word/document.xml", "xlsx": "xl/workbook.xml", "pptx": "ppt/presentation.xml"}
            with zipfile.ZipFile(path) as archive:
                missing = not {"[Content_Types].xml", required[converter.id]} <= set(archive.namelist())
            if missing or isinstance(exc, KeyError):
                raise ConversionError("missing_part", str(exc)) from exc
            if isinstance(exc, ValueError):
                raise ConversionError("malformed", str(exc)) from exc
        if converter.id == "image" and isinstance(exc, (EOFError, SyntaxError)):
            raise ConversionError("malformed", str(exc)) from exc
        if isinstance(exc, csv.Error):
            code = "resource_limit" if str(exc).startswith("field larger than field limit") else "malformed"
            raise ConversionError(code, str(exc)) from exc
        raise
    # PDF already emits image_semantics_unverified itself, as before.
    if converter.id == "pdf":
        return
    for page in document.pages:
        for block in page.blocks:
            if block.kind == "image":
                page.warnings.append("image_semantics_unverified")
