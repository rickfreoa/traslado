# SPDX-License-Identifier: AGPL-3.0-or-later
# Copyright (C) 2026 Ricardo Peres Freoa
"""Residual E3/E4 integration regressions."""
import json
from pathlib import Path
import subprocess
import sys
from types import SimpleNamespace as N
import unittest
from unittest.mock import patch

try:
    from tests.test_legacy_e4 import scratch
except ImportError:  # suíte rodando de cópia dos testes, sem o pacote tests
    from test_legacy_e4 import scratch
from traslado.errors import ConversionError, classify, partial_error
from traslado.model import Document, Options
from traslado.runner import discover, execute


class ResidualIntegration(unittest.TestCase):
    def test_parser_exception_classes(self):
        import fitz
        from lxml.etree import XMLSyntaxError
        from PIL import UnidentifiedImageError
        from docx.oxml.exceptions import InvalidXmlError as DocxInvalid
        from pptx.exc import InvalidXmlError as PptxInvalid
        for exc in (fitz.FileDataError('bad'), fitz.EmptyFileError('empty'),
                    XMLSyntaxError('bad', 0, 1, 1), UnidentifiedImageError('bad'),
                    DocxInvalid('bad'), PptxInvalid('bad')):
            with self.subTest(kind=type(exc)):
                self.assertEqual(classify(exc).code, 'malformed')
        self.assertEqual(classify(RuntimeError('bug')).code, 'internal')

    def test_docx_non_numeric_grid_attributes(self):
        from docx import Document as Docx
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        with scratch() as root:
            for kind in ('gridSpan', 'gridBefore'):
                with self.subTest(kind=kind):
                    doc = Docx()
                    table = doc.add_table(rows=1, cols=1)
                    table.cell(0, 0).text = 'sentinel'
                    parent = (table.cell(0, 0)._tc.get_or_add_tcPr() if kind == 'gridSpan'
                              else table.rows[0]._tr.get_or_add_trPr())
                    node = OxmlElement('w:' + kind)
                    node.set(qn('w:val'), 'abc')
                    parent.append(node)
                    source = root / (kind + '.docx')
                    doc.save(source)
                    result = execute(source, root / 'out', Options(mode='native'))
                    self.assertEqual(result.get('error_code'), 'malformed', result)

    def test_pdf_localized_subclasses_and_table_failures(self):
        import fitz
        from traslado.pdf import extract
        from traslado import pdf as pdf_module
        from PIL import Image
        class MuPdfFailure(RuntimeError):
            pass
        with scratch() as root:
            source = root / 'table.pdf'
            image = root / 'image.png'
            Image.new('RGB', (40, 40), 'white').save(image)
            with fitz.open() as pdf:
                page = pdf.new_page()
                page.insert_text((40, 40), 'Native sentinel text 12345')
                page.insert_image(fitz.Rect(40, 150, 180, 290), filename=str(image))
                for y in (80, 100, 120):
                    page.draw_line((40, y), (200, y))
                pdf.save(source)
            for method, exc in [('render_page', MuPdfFailure('bad')),
                                ('render_page', fitz.FileDataError('bad')),
                                ('render_page', ConversionError('resource_limit', 'limit')),
                                *[('find_tables', exc) for exc in
                                  (MuPdfFailure('bad'), fitz.FileDataError('bad'),
                                   ValueError('bad'), IndexError('bad'), ZeroDivisionError('bad'),
                                   ConversionError('resource_limit', 'limit'))]]:
                with self.subTest(method=method, kind=type(exc)):
                    doc = Document(source.name, 'hash')
                    with patch.object(pdf_module if method == 'render_page' else fitz.Page, method, side_effect=exc):
                        extract(source, root, Options(mode='native'), doc)
                    self.assertEqual(partial_error(doc).code,
                                     'resource_limit' if isinstance(exc, ConversionError) else 'malformed')

    def test_mixed_directory_opt_in_and_cli_exit(self):
        with scratch() as root:
            inputs = root / 'inputs'
            inputs.mkdir()
            (inputs / 'yes.txt').write_text('sentinel', encoding='utf-8')
            (inputs / 'review.html').write_text('<p>sentinel</p>', encoding='utf-8')
            (inputs / 'legacy.rtf').write_bytes(b'{\\rtf1 sentinel}')
            self.assertEqual({p.suffix for p in discover(inputs, root / 'out')}, {'.txt', '.html'})
            self.assertEqual(len(discover(inputs, root / 'out', legacy=True)), 3)
            self.assertEqual(discover(inputs / 'legacy.rtf', root / 'out'), [inputs / 'legacy.rtf'])
            for partial, expected in ((False, 0), (True, 3)):
                if partial:
                    (inputs / 'review.ignored').rename(inputs / 'review.html')
                else:
                    (inputs / 'review.html').rename(inputs / 'review.ignored')
                proc = subprocess.run([sys.executable, '-X', 'utf8', '-m', 'traslado', 'convert',
                                       str(inputs), '-o', str(root / ('out' + str(partial)))],
                                      capture_output=True, text=True, encoding='utf-8', timeout=90)
                self.assertEqual(proc.returncode, expected, proc.stdout + proc.stderr)
                self.assertNotIn('legacy.rtf', proc.stdout)
            result = execute(inputs / 'legacy.rtf', root / 'explicit')
            self.assertEqual(result.get('error_code'), 'unsupported')
            self.assertIn('--engine legacy', result['error'])

    def test_legacy_images_rebased_and_code_protected(self):
        from traslado.legacy import _map
        from traslado.render import render
        from traslado.chunks import write_chunks
        text = lambda value: N(kind='text', text=value)
        para = lambda items: N(kind='paragraph', content=items)
        literal = '![Figura](assets/legacy-1.png)'
        source = N(notes=[], assets=[N(id=0, media_type='image/png', origin_part='img', data=b'image bytes')],
                   blocks=[para([text('Antes '), N(kind='image', alt='Figura', source=N(kind='asset', asset_id=0)), text(' depois')]),
                           para([text('`' + literal + '`')]),
                           N(kind='code_block', text=literal, lang='md')])
        with scratch() as root:
            doc = Document('image.rtf', 'hash')
            _map(source, root, Options(images=True), doc)
            md = render(doc)
            write_chunks(doc, md, root / 'chunks')
            chunks = ''.join(p.read_text(encoding='utf-8') for p in sorted((root / 'chunks').glob('*.md')))
            self.assertIn('Antes ![Figura](../assets/legacy-1.png) depois', chunks)
            self.assertIn('`' + literal + '`', chunks)
            self.assertIn('```md\n' + literal + '\n```', chunks)
            self.assertEqual((root / 'assets/legacy-1.png').read_bytes(), b'image bytes')
            entries = [json.loads(line) for line in (root / 'chunks/chunks.jsonl').read_text(encoding='utf-8').splitlines()]
            for entry in entries:
                self.assertEqual(entry['text'], md[entry['char_start']:entry['char_end']])

    def test_note_leading_space_only(self):
        from traslado.legacy import _map
        para = lambda value: N(kind='paragraph', content=[N(kind='text', text=value)])
        source = N(assets=[], blocks=[N(kind='paragraph', content=[N(kind='note_ref', note_id='n')])],
                   notes=[N(id='n', kind='footnote', blocks=[para('  Nota  interna  ')])])
        with scratch() as root:
            doc = Document('note.rtf', 'hash')
            _map(source, root, Options(), doc)
            self.assertEqual(doc.pages[0].blocks[-1].text, '[^1]: Nota  interna  ')

    def test_legacy_warning_labels(self):
        from traslado.review import LABELS, warning_label
        for code in ('legacy_engine_experimental', 'legacy_layout_unavailable',
                     'legacy_table_merged_cells_need_review', 'legacy_note_definition_unreferenced',
                     'legacy_image_unavailable'):
            self.assertIn(code, LABELS)
            label = warning_label(code)
            self.assertTrue(label.endswith(' — ' + code))
            self.assertNotEqual(label.split(' — ')[0], code.replace('_', ' '))

    def test_real_rtf_image_fixture_chunks(self):
        import base64
        png = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII=')
        with scratch() as root:
            source = root / 'image.rtf'
            source.write_text(r'{\rtf1\ansi Before {\pict\pngblip ' + png.hex() + r'} after}', encoding='ascii')
            result = execute(source, root / 'out', Options(legacy=True, mode='native', chunks=True))
            self.assertNotIn('error_code', result, result)
            output = Path(result['output'])
            self.assertEqual((output / 'assets/legacy-1.png').read_bytes(), png)
            md = (output / 'document.md').read_text(encoding='utf-8')
            self.assertIn('](assets/legacy-1.png)', md)
            chunks = ''.join(p.read_text(encoding='utf-8') for p in (output / 'chunks').glob('*.md'))
            self.assertIn('](../assets/legacy-1.png)', chunks)
