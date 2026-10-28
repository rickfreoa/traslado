"""T11: registry contract and published-byte parity with the frozen 1.0.0 tree."""
from __future__ import annotations

from dataclasses import FrozenInstanceError, fields, replace
from contextlib import contextmanager
import ast
import builtins
import json
import os
from pathlib import Path
import subprocess
import sys
import shutil
import uuid
import unittest
from unittest.mock import patch
import zipfile
import io

from traslado import doctor, formats, registry, runner
from traslado.model import Document, Options


EXPECTED = frozenset(".pdf .docx .xlsx .pptx .txt .md .json .xml .csv .tsv .html .htm .png .jpg .jpeg .tif .tiff .bmp .webp .doc .rtf .odt .ods .odp .ppt .epub".split())
BUILTIN_EXPECTED = EXPECTED - frozenset(".doc .rtf .odt .ods .odp .ppt .epub".split())
ROOT = Path(__file__).resolve().parents[1]


def _has_v100_tag() -> bool:
    try:
        return subprocess.run(["git", "rev-parse", "-q", "--verify", "refs/tags/v1.0.0"], cwd=ROOT,
                              capture_output=True).returncode == 0
    except OSError:
        return False


HAS_V100 = _has_v100_tag()
import importlib.util  # noqa: E402
FULL = all(importlib.util.find_spec(m) for m in ("fitz", "PIL", "docx", "openpyxl", "pptx"))


@contextmanager
def scratch():
    # Python 3.14's Windows mode-0700 temp ACL is inaccessible to the sandbox
    # token. Ordinary test directories inherit the authorized workspace ACL.
    root = ROOT / ".t11-logs" / ("registry-" + uuid.uuid4().hex)
    root.mkdir(parents=True)
    try:
        yield str(root)
    finally:
        shutil.rmtree(root)


class RegistryT11(unittest.TestCase):
    def test_inventory_is_unique_frozen_and_compatible(self):
        self.assertEqual(registry.SUPPORTED, EXPECTED)
        self.assertIs(formats.SUPPORTED, registry.SUPPORTED)
        self.assertIs(runner.SUPPORTED, registry.SUPPORTED)
        self.assertIsInstance(registry.REGISTRY, tuple)
        self.assertEqual({c.id for c in registry.REGISTRY},
                         {"pdf", "docx", "xlsx", "pptx", "image", "csv", "html", "text", "legacy"})
        self.assertEqual(sum(len(c.extensions) for c in registry.REGISTRY), len(EXPECTED))
        for converter in registry.REGISTRY:
            with self.assertRaises(FrozenInstanceError):
                converter.id = "changed"
            for name in ("extensions", "capabilities", "limits", "known_losses", "requires", "optional"):
                self.assertIsInstance(getattr(converter, name), tuple)
            self.assertLessEqual(set(converter.limits), {f.name for f in fields(Options)})
            for suffix in converter.extensions:
                self.assertEqual(suffix, suffix.lower())
                self.assertTrue(suffix.startswith("."))
                self.assertIs(registry.lookup(suffix.upper()), converter)
        self.assertIsNone(registry.lookup(".unsupported"))

    def test_status_codes_and_dependency_precedence(self):
        with patch.object(registry.importlib.util, "find_spec", return_value=object()), \
                patch.object(registry.ocr, "find_tesseract", return_value="tesseract"):
            for converter in registry.REGISTRY:
                self.assertEqual(converter.status().state, "available")
                self.assertEqual(converter.status().code, "ok")
        with patch.object(registry.importlib.util, "find_spec", return_value=object()), \
                patch.object(registry.ocr, "find_tesseract", return_value=None):
            for name in ("pdf", "image"):
                converter = next(c for c in registry.REGISTRY if c.id == name)
                self.assertEqual((converter.status().state, converter.status().code),
                                 ("degraded", "missing_engine:tesseract"))
        for converter in registry.REGISTRY:
            for dependency in (*converter.requires, *converter.optional):
                with self.subTest(converter=converter.id, dependency=dependency), \
                        patch.object(registry.importlib.util, "find_spec",
                                     side_effect=lambda name: None if name == dependency else object()), \
                        patch.object(registry.ocr, "find_tesseract", return_value=None):
                    status = converter.status()
                    self.assertIn(status.code, registry.REASON_CODES)
                    self.assertEqual(status.state, "unavailable" if dependency in converter.requires else "degraded")
                    if dependency in converter.requires:
                        self.assertEqual(status.code, "missing_module:" + dependency)

    def test_status_does_not_import_heavy_libraries(self):
        code = """
import sys
from traslado.registry import REGISTRY
for converter in REGISTRY:
    converter.status()
assert not {'fitz', 'PIL', 'docx', 'openpyxl', 'pptx'} & set(sys.modules)
"""
        completed = subprocess.run([sys.executable, "-X", "utf8", "-c", code],
                                   capture_output=True, text=True)
        self.assertEqual(completed.returncode, 0, completed.stderr)

    def test_discover_execute_and_doctor_agree_on_every_extension(self):
        with scratch() as folder:
            root = Path(folder)
            inputs = root / "inputs"
            inputs.mkdir()
            for suffix in EXPECTED | {".unsupported"}:
                (inputs / ("record" + suffix.upper())).write_bytes(b"fixture")
            with patch.object(doctor, "find_tesseract", return_value=None), \
                    patch.object(registry.ocr, "find_tesseract", return_value=None):
                report = doctor.inspect()
            self.assertEqual([c["id"] for c in report["converters"]], sorted(c.id for c in registry.REGISTRY))
            listed = {ext for item in report["converters"] for ext in item["extensions"]}
            self.assertEqual(listed, EXPECTED)
            for item in report["converters"]:
                self.assertEqual(set(item), {"id", "extensions", "state", "code", "detail", "capabilities", "known_losses"})
                self.assertIn(item["code"], registry.REASON_CODES)
            self.assertEqual({p.suffix.lower() for p in runner.discover(inputs, root / "out")}, BUILTIN_EXPECTED)
            self.assertEqual({p.suffix.lower() for p in runner.discover(inputs, root / "out", legacy=True)}, EXPECTED)
            for source in inputs.iterdir():
                supported = source.suffix.lower() in EXPECTED
                # Single files keep the 1.0.0 path: execute refuses them with the established message.
                self.assertEqual(runner.discover(source, root / "out"), [source])
                # E3 must bound input before reading signatures, even for an
                # unknown extension (whose content might be supported).
                result = runner.execute(source, root / "out", Options(max_input_bytes=1))
                self.assertEqual(result["error"], "Entrada excede limite de bytes")
                self.assertEqual(result["error_code"], "resource_limit")
                if not supported:
                    result = runner.execute(source, root / "out")
                    self.assertEqual(result["error_code"], "unsupported")
                    self.assertEqual(result["error"],
                                     "Formato não suportado; use PDF, DOCX, XLSX, PPTX, texto, CSV, HTML ou imagem")
            # Directory traversal and single-file discovery use the same inventory.
            nested = inputs / "nested"
            nested.mkdir()
            (nested / "nested.TXT").write_bytes(b"nested")
            self.assertEqual(len(runner.discover(inputs, root / "out", recursive=True)), len(BUILTIN_EXPECTED) + 1)

    def test_unavailable_does_not_short_circuit_worker(self):
        with scratch() as folder:
            root = Path(folder)
            source = root / "broken.pdf"
            source.write_bytes(b"%PDF broken")
            original_import = builtins.__import__
            def without_fitz(name, *args, **kwargs):
                if name == "fitz":
                    raise ModuleNotFoundError("No module named 'fitz'")
                return original_import(name, *args, **kwargs)
            with patch.object(registry.Converter, "status", side_effect=AssertionError("must not probe")), \
                    patch.object(builtins, "__import__", side_effect=without_fitz):
                with self.assertRaisesRegex(RuntimeError, r"Instale o perfil PDF: pip install \.\[pdf\]"):
                    formats.extract(source, root, Options(), Document(source.name, "hash"))

    @unittest.skipUnless(HAS_V100, "requer checkout git com a tag v1.0.0")
    def test_original_branch_bodies_are_unchanged(self):
        class BeforeE3(ast.NodeTransformer):
            """Ignore only E3's typed failures and localized-error bookkeeping.

            Published-byte parity below remains the behavioral preservation gate.
            """
            def visit_Raise(self, node):
                self.generic_visit(node)
                if isinstance(node.exc, ast.Call) and isinstance(node.exc.func, ast.Name) and node.exc.func.id == 'ConversionError':
                    node.exc = ast.Call(func=ast.Name(id='ValueError', ctx=ast.Load()),
                                        args=node.exc.args[1:], keywords=[])
                return node

            def visit_Expr(self, node):
                if isinstance(node.value, ast.Call) and isinstance(node.value.func, ast.Name) and node.value.func.id in {'localized', '_load_image'}:
                    return None
                return self.generic_visit(node)

            def visit_Call(self, node):
                self.generic_visit(node)
                if isinstance(node.func, ast.Name) and node.func.id == '_check_ooxml':
                    return ast.Call(func=ast.Name(id='check_archive', ctx=ast.Load()), args=node.args[:2], keywords=[])
                if isinstance(node.func, ast.Name) and node.func.id == '_parser_iter':
                    return node.args[0]
                if isinstance(node.func, ast.Name) and node.func.id == '_parser_call':
                    call = node.args[0]
                    return call.body if isinstance(call, ast.Lambda) else ast.Call(func=call, args=node.args[1:], keywords=node.keywords)
                if isinstance(node.func, ast.Name) and node.func.id == 'exception_name':
                    return ast.Attribute(value=ast.Call(func=ast.Name(id='type', ctx=ast.Load()),
                                                       args=node.args, keywords=[]),
                                         attr='__name__', ctx=ast.Load())
                return node

        original = ast.parse(subprocess.check_output(
            ["git", "show", "v1.0.0:src/freoa_v6/formats.py"], cwd=ROOT).decode("utf-8"))
        current = BeforeE3().visit(ast.parse(Path(formats.__file__).read_text(encoding="utf-8")))
        extract = next(n for n in original.body if isinstance(n, ast.FunctionDef) and n.name == "extract")
        branch = next(n for n in extract.body if isinstance(n, ast.If) and ast.unparse(n.test) == "suffix == '.docx'")
        for name in ("docx", "xlsx", "pptx", "image"):
            function = next(n for n in current.body if isinstance(n, ast.FunctionDef) and n.name == "_extract_" + name)
            if name == "docx":  # A5a rewrites the DOCX branch on purpose; its contract has its own tests.
                branch = branch.orelse[0] if len(branch.orelse) == 1 else branch.orelse
                continue
            self.assertEqual([ast.dump(n) for n in branch.body],
                             [ast.dump(n) for n in function.body[-len(branch.body):]], name)
            branch = branch.orelse[0] if len(branch.orelse) == 1 else branch.orelse

    def test_dispatch_uses_registered_callable_and_preserves_image_postprocess(self):
        def extractor(path, workspace, options, document):
            from traslado.model import Block, Page
            document.pages.append(Page(1, "native", [Block("image", asset="assets/test.png")]))
        converter = replace(registry.lookup(".txt"), extract=extractor)
        with scratch() as folder, patch.object(formats, "lookup", return_value=converter) as lookup:
            document = Document("test.txt", "hash")
            formats.extract(Path("test.txt"), Path(folder), Options(), document)
            lookup.assert_called_once_with(".txt")
            self.assertEqual(document.pages[0].warnings, ["image_semantics_unverified"])

    @unittest.skipUnless(HAS_V100 and FULL, "requer checkout git com a tag v1.0.0 e o perfil Full")
    def test_published_bytes_match_v100_for_all_extensions(self):
        """Generate one fixture per extension, then run both real worker pipelines."""
        from PIL import Image
        import fitz
        from docx import Document as Docx
        from openpyxl import Workbook
        from pptx import Presentation
        from pptx.util import Inches

        with scratch() as folder:
            root = Path(folder)
            baseline = root / "baseline"
            archive = subprocess.check_output(["git", "archive", "--format=zip", "v1.0.0"], cwd=ROOT)
            with zipfile.ZipFile(io.BytesIO(archive)) as package:
                package.extractall(baseline)
            inputs = root / "inputs"
            inputs.mkdir()
            image = inputs / "record.png"
            Image.new("RGB", (120, 80), "white").save(image)
            for suffix in (".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"):
                with Image.open(image) as im:
                    im.save(inputs / ("record" + suffix))
            pdf = fitz.open()
            page = pdf.new_page()
            page.insert_text((40, 40), "NATIVE_SENTINEL 12345")
            page.insert_image(fitz.Rect(40, 80, 160, 160), filename=str(image))
            pdf.save(inputs / "record.pdf")
            pdf.close()
            doc = Docx()
            doc.add_heading("Heading", level=1)
            doc.add_paragraph("BEFORE_SENTINEL 12345")
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).merge(table.cell(0, 1)).text = "MERGED 98765"
            table.cell(1, 0).text = "Value"
            table.cell(1, 1).text = "456"
            doc.add_picture(str(image))
            doc.save(inputs / "record.docx")
            wb = Workbook()
            wb.active.append(["Value", "Formula"])
            wb.active.append([12345, "=A2*2"])
            wb.create_sheet("Second").append(["SECOND_SENTINEL"])
            wb.save(inputs / "record.xlsx")
            wb.close()
            prs = Presentation()
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "SLIDE_SENTINEL 12345"
            slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table.cell(0, 0).text = "TABLE 98765"
            slide.shapes.add_picture(str(image), Inches(1), Inches(4))
            prs.save(inputs / "record.pptx")
            for suffix, content in {
                ".txt": "BEFORE\nConteúdo 12345\nAFTER\n", ".md": "# Title\n\nBody `code` 12345\n",
                ".json": '{"value": "``` 12345"}', ".xml": "<value>12345</value>",
                ".csv": 'name,note\nAlice,"first\nsecond"\n', ".tsv": "name\tnote\nAlice\t12345\n",
                ".html": "<h1>Title</h1><table><tr><td>12345</td></tr></table><script>hidden</script>",
                ".htm": "<p>BEFORE &amp; AFTER 12345</p>",
            }.items():
                (inputs / ("record" + suffix)).write_text(content, encoding="utf-8")
            self.assertEqual({p.suffix for p in inputs.iterdir()}, BUILTIN_EXPECTED)
            code = """
import json,sys,hashlib,subprocess
from dataclasses import asdict
from pathlib import Path
from traslado.model import Options
inputs, output = map(Path, sys.argv[1:3])
results = {}
for source in sorted(inputs.iterdir()):
    workspace = output / source.suffix[1:]
    workspace.mkdir(parents=True)
    request = workspace.parent / (source.suffix[1:] + '-request.json')
    request.write_text(json.dumps({'source_name': source.name,
        'sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
        'snapshot': str(source), 'workspace': str(workspace),
        # furniture="off" (1.1 only) reproduces the 1.0.0 state; the 1.0.0 Options has no such field.
        'options': asdict(Options(mode='native', chunks=True,
                                  **({'furniture': 'off'} if 'furniture' in Options.__dataclass_fields__ else {})))}),
        encoding='utf-8')
    proc = subprocess.run([sys.executable, '-P', '-m', 'traslado.worker', str(request)],
                          capture_output=True, text=True)
    assert proc.returncode == 0, (source, proc.stderr, (workspace/'error.json').read_text() if (workspace/'error.json').exists() else '')
    results[source.suffix] = str(workspace)
print(json.dumps(results))
"""
            outputs = []
            for tree, label in ((baseline, "baseline"), (ROOT, "candidate")):
                env = os.environ.copy()
                env["PYTHONPATH"] = str(tree / "src")
                program = code.replace("traslado", "freoa_v6") if label == "baseline" else code
                proc = subprocess.run([sys.executable, "-X", "utf8", "-c", program, str(inputs), str(root / label / "out")],
                                      env=env, capture_output=True, text=True, encoding="utf-8", timeout=180)
                self.assertEqual(proc.returncode, 0, proc.stderr)
                outputs.append(json.loads(proc.stdout))
            compared = []
            # DOCX changes on purpose in A5a (notes, per-species warnings) and has its own tests.
            for suffix in sorted(BUILTIN_EXPECTED - {".docx"}):
                for relative in ("document.md", "document.json", "document-layout.json", "chunks/chunks.jsonl"):
                    payloads = [(Path(out[suffix]) / relative).read_bytes() for out in outputs]
                    if relative == "document.json":
                        # Neutralize just these fields (timing; release version; options added in 1.1); every other
                        # field must be equal.
                        def neutral(data):
                            data = json.loads(data)
                            data.update(duration_seconds=None, converter_version=None)
                            data["options"] = {k: v for k, v in data.get("options", {}).items()
                                               if k not in ("ocr_workers", "furniture")}
                            return data
                        payloads = [neutral(data) for data in payloads]
                    self.assertEqual(payloads[0], payloads[1], (suffix, relative))
                    compared.append(suffix + ":" + relative)
            logs = ROOT / ".t11-logs"
            logs.mkdir(exist_ok=True)
            (logs / "byte-parity.json").write_text(json.dumps({"baseline": "v1.0.0", "compared": compared,
                                                             "equal": len(compared)}, indent=2), encoding="utf-8")


if __name__ == "__main__":
    unittest.main()
