from __future__ import annotations

import io
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

from freoa_v6.advanced import sanitize
from freoa_v6 import formats
from freoa_v6.model import Document, Options
from freoa_v6.runner import execute

try:
    from markdown_it import MarkdownIt
except ImportError:
    for parent in Path(__file__).resolve().parents:
        oracle = parent / 'work/review-next/md-parser'
        if oracle.is_dir():
            sys.path.insert(0, str(oracle))
            break
    try:
        from markdown_it import MarkdownIt
    except ImportError:
        MarkdownIt = None

def presentation(path, missing_coordinates=False, image=False):
    from PIL import Image
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    first = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1))
    first.text = 'FIRST_SOURCE_SENTINEL'
    second = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1))
    second.text = 'SECOND_SOURCE_SENTINEL'
    table = slide.shapes.add_table(1, 1, Inches(1), Inches(3), Inches(2), Inches(1)).table
    table.cell(0, 0).text = 'TABLE_SOURCE_SENTINEL'
    if missing_coordinates:
        first._element.spPr.remove(first._element.spPr.xfrm)
    if image:
        data = io.BytesIO()
        Image.new('RGB', (80, 80), 'white').save(data, 'PNG')
        data.seek(0)
        slide.shapes.add_picture(data, Inches(1), Inches(4), width=Inches(1))
    prs.save(path)

class MinorRC3(unittest.TestCase):
    def test_self_closing_safe_and_unknown_tags_do_not_invent_end_tags(self):
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw)
            content = '<table><tr><td>A<br/>B<br />C</td><td colspan="2" onclick="evil()"/></tr></table><img src="https://example.invalid/a"/>'
            safe = sanitize(content, root, root, 1, Options())
            self.assertIn('A<br>B<br>C', safe)
            self.assertIn('<td colspan="2"/>', safe)
            self.assertNotIn('onclick', safe)
            self.assertNotIn('&lt;/br&gt;', safe)
            self.assertNotIn('&lt;/img&gt;', safe)
            self.assertIn('&lt;img src=&quot;https://example.invalid/a&quot;/&gt;', safe)

    @unittest.skipIf(MarkdownIt is None, 'Optional CommonMark oracle unavailable')
    def test_selfclosing_html_keeps_model_link_and_image_guards(self):
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw)
            content = '<br/><img src="https://example.invalid/a"/>\n![a\nb](https://example.invalid/a.png)\n> ![ref]\n>\n> [ref]: https://example.invalid/a.png'
            safe = sanitize(content, root, root, 1, Options())
            tokens = MarkdownIt('commonmark', {'html': True}).parse(safe)
            def links(nodes):
                for node in nodes:
                    if node.type in {'image', 'link_open'}:
                        yield node
                    if node.children:
                        yield from links(node.children)
            self.assertEqual(list(links(tokens)), [])

    def test_real_pptx_missing_coordinates_preserves_text_table_and_document_order(self):
        from pptx import Presentation
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw)
            source = root / 'coordinates.pptx'
            presentation(source, missing_coordinates=True)
            self.assertIsNone(Presentation(source).slides[0].shapes[0].top)
            result = execute(source, root / 'out', Options())
            self.assertEqual(result['status'], 'partial', result)
            package = Path(result['output'])
            markdown = (package / 'document.md').read_text(encoding='utf-8')
            for sentinel in ['FIRST_SOURCE_SENTINEL', 'SECOND_SOURCE_SENTINEL', 'TABLE_SOURCE_SENTINEL']:
                self.assertIn(sentinel, markdown)
            self.assertLess(markdown.index('FIRST_SOURCE_SENTINEL'), markdown.index('SECOND_SOURCE_SENTINEL'))
            doc = json.loads((package / 'document.json').read_text(encoding='utf-8'))
            self.assertIn('pptx_shape_coordinates_missing_document_order', doc['pages'][0]['warnings'])

    def test_real_pptx_image_pixel_limit_is_partial_and_preserves_other_blocks(self):
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw)
            source = root / 'image-limit.pptx'
            presentation(source, image=True)
            result = execute(source, root / 'out', Options(max_pixels=1000))
            self.assertEqual(result['status'], 'partial', result)
            package = Path(result['output'])
            doc = json.loads((package / 'document.json').read_text(encoding='utf-8'))
            self.assertTrue(any(w.startswith('pptx_image_failed:') and 'ValueError' in w for w in doc['pages'][0]['warnings']))
            markdown = (package / 'document.md').read_text(encoding='utf-8')
            self.assertIn('FIRST_SOURCE_SENTINEL', markdown)
            self.assertIn('TABLE_SOURCE_SENTINEL', markdown)
            self.assertFalse(any(b['kind'] == 'image' for b in doc['pages'][0]['blocks']))
            self.assertEqual(list((package / 'assets').iterdir()), [])

    def test_failed_image_partial_asset_is_removed_without_losing_slide_blocks(self):
        with tempfile.TemporaryDirectory() as raw:
            root = Path(raw)
            source = root / 'partial-image.pptx'
            presentation(source, image=True)
            workspace = root / 'package'
            workspace.mkdir()
            def partial_image(data, target, options):
                target.write_bytes(b'INCOMPLETE_IMAGE_SENTINEL')
                raise OSError('synthetic interrupted image save')
            doc = Document(source.name, '0' * 64)
            with patch.object(formats, 'image_to_png', partial_image):
                formats.extract(source, workspace, Options(), doc)
            self.assertEqual(list((workspace / 'assets').iterdir()), [])
            self.assertTrue(any('OSError' in w for w in doc.pages[0].warnings))
            self.assertTrue(any(b.kind == 'table' for b in doc.pages[0].blocks))
            self.assertTrue(any(b.text == 'FIRST_SOURCE_SENTINEL' for b in doc.pages[0].blocks))

if __name__ == '__main__':
    unittest.main()
