# SPDX-License-Identifier: AGPL-3.0-or-later
# Copyright (C) 2026 Ricardo Peres Freoa
from pptx import Presentation
p = Presentation(); s = p.slides.add_slide(p.slide_layouts[1])
s.shapes.title.text = 'Relato publico'
tf = s.placeholders[1].text_frame; tf.text = 'Acao inicial.'; tf.add_paragraph().text = 'Fim do relato.'
p.save(__import__('os').path.join(__import__('os').path.dirname(__file__), 'ppt-fonte.pptx'))
