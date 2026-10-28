# SPDX-License-Identifier: AGPL-3.0-or-later
"""Independent synthetic acceptance oracles; no customer documents or downloads."""
from __future__ import annotations

import hashlib
import io
import json
import os
from pathlib import Path
import stat
import subprocess
import sys
import tempfile
import time
import unittest
import zipfile

from traslado.model import Block, Options
from traslado.pdf import reading_order
from traslado.safety import check_archive, run_bounded, Cancelled
from traslado.runner import execute


class Acceptance(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="freoa-acceptance-")
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def text(self, name, text):
        path = self.root / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(text, encoding="utf-8")
        return path

    def job(self, source, options=None, output=None):
        return execute(source, output or self.root / "results", options or Options())

    def artifact(self, result):
        self.assertIn(result["status"], {"complete", "partial"})
        root = Path(result["output"])
        self.assertTrue(root.is_dir())
        self.assertTrue((root / "manifest.json").is_file())
        manifest=json.loads((root/"manifest.json").read_text(encoding="utf-8"))
        for member in manifest["files"]:
            path=(root/member["path"]).resolve()
            self.assertTrue(path.is_relative_to(root.resolve()))
            self.assertEqual(path.stat().st_size,member["bytes"])
            self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(),member["sha256"])
        document = json.loads((root / "document.json").read_text(encoding="utf-8"))
        md = (root / "document.md").read_text(encoding="utf-8")
        return root, document, md

    def pdf(self, name, lines, width=600, height=800):
        import fitz
        document = fitz.open()
        for page_lines in lines:
            page = document.new_page(width=width, height=height)
            for index, text in enumerate(page_lines):
                page.insert_text((50, 70 + index*25), text, fontsize=12)
        target = self.root / name
        document.save(target)
        document.close()
        return target

    def scan(self):
        import fitz
        from PIL import Image, ImageDraw, ImageFont
        image = Image.new("RGB", (1600, 900), "white")
        font_path = Path(os.environ.get("SystemRoot", "C:/Windows")) / "Fonts/arial.ttf"
        try:
            font = ImageFont.truetype(str(font_path), 42)
        except OSError:
            font = ImageFont.truetype("DejaVuSans.ttf", 42)
        draw = ImageDraw.Draw(image)
        expected = ["CONTRATO DE SERVICOS", "Valor total: 12345 reais.", "Prazo: 30 dias."]
        for index, text in enumerate(expected):
            draw.text((100, 120+index*150), text, font=font, fill="black")
        stream=io.BytesIO();image.save(stream, "PNG")
        document=fitz.open();page=document.new_page(width=800,height=450)
        page.insert_image(page.rect,stream=stream.getvalue())
        target=self.root/"scan.pdf";document.save(target);document.close()
        return target, expected

    def test_options_reject_boolean_limits_and_injection_language(self):
        for values in ({"job_timeout": True}, {"dpi": 0}, {"language": "eng --psm 0"}, {"tables": "false"}):
            with self.subTest(values=values), self.assertRaises(ValueError):
                Options(**values)

    def test_plain_text_sentinel_and_source_immutability(self):
        source=self.text("record.txt", "BEFORE_SENTINEL\nValor: 456789\nAFTER_SENTINEL")
        original=source.read_bytes()
        result=self.job(source)
        root,document,md=self.artifact(result)
        self.assertEqual(result["status"], "complete")
        self.assertEqual(source.read_bytes(), original)
        self.assertIn("Valor: 456789",md)
        self.assertLess(md.index("BEFORE_SENTINEL"),md.index("AFTER_SENTINEL"))
        self.assertEqual(document["source_sha256"],hashlib.sha256(original).hexdigest())

    def test_same_source_and_output_cannot_mutate_original(self):
        source=self.text("source.md", "# ORIGINAL\n\nORIGINAL_SENTINEL")
        original=source.read_bytes()
        try:
            result=self.job(source,output=source)
        except (ValueError,OSError):
            pass
        else:
            self.assertEqual(result["status"],"failed")
        self.assertEqual(source.read_bytes(),original)

    def test_same_name_distinct_inputs_retain_both_contents(self):
        first=self.text("a/contract.txt","FIRST_DOCUMENT_ONLY")
        second=self.text("b/contract.txt","SECOND_DOCUMENT_ONLY")
        one=self.job(first);two=self.job(second)
        r1,d1,m1=self.artifact(one);r2,d2,m2=self.artifact(two)
        self.assertNotEqual(r1,r2)
        self.assertIn("FIRST_DOCUMENT_ONLY",m1)
        self.assertIn("SECOND_DOCUMENT_ONLY",m2)

    def test_changed_input_retains_predecessor_package(self):
        source=self.text("contract.txt","FIRST_IMMUTABLE_RECORD")
        first=self.job(source);root,doc,md=self.artifact(first)
        snapshot={p.relative_to(root).as_posix():p.read_bytes() for p in root.rglob("*") if p.is_file()}
        source.write_text("SECOND_RECORD",encoding="utf-8")
        second=self.job(source);root2,doc2,md2=self.artifact(second)
        self.assertNotEqual(root,root2)
        self.assertIn("SECOND_RECORD",md2)
        self.assertEqual(snapshot,{p.relative_to(root).as_posix():p.read_bytes() for p in root.rglob("*") if p.is_file()})

    def test_csv_multiline_value_preserved_in_structured_output(self):
        source=self.text("record.csv",'nome,observacao\nAlice,"primeira\nsegunda"\n')
        _,document,md=self.artifact(self.job(source))
        rows=[b["rows"] for p in document["pages"] for b in p["blocks"] if b["kind"]=="table"]
        self.assertIn(["Alice","primeira\nsegunda"], rows[0])
        self.assertNotIn("primeirasegunda",md)

    def test_html_table_numbers_survive_and_loss_is_explicit(self):
        source=self.text("record.html",'<p>INTRO_SENTINEL</p><table><tr><td>Objeto</td><td>98765</td></tr></table><script>SECRET_SCRIPT</script>')
        result=self.job(source);_,document,md=self.artifact(result)
        self.assertIn("98765",md)
        self.assertIn("INTRO_SENTINEL",md)
        self.assertNotIn("SECRET_SCRIPT",md)
        self.assertTrue(any("html_table" in w for p in document["pages"] for w in p["warnings"]))
        self.assertEqual(result["status"],"partial")

    def test_short_regular_pdf_uses_native_without_ocr_dependency(self):
        source=self.pdf("short.pdf",[["SHORT_NATIVE_SENTINEL 12345"]])
        result=self.job(source,Options(tesseract=str(self.root/"missing.exe")))
        _,document,md=self.artifact(result)
        self.assertEqual(result["status"],"complete")
        self.assertEqual(document["pages"][0]["route"],"native")
        self.assertIn("SHORT_NATIVE_SENTINEL 12345",md)
        self.assertFalse(any(b.get("engine")=="tesseract" for p in document["pages"] for b in p["blocks"]))

    def test_blank_pdf_page_does_not_require_ocr(self):
        source=self.pdf("blank.pdf",[[]])
        result=self.job(source,Options(tesseract=str(self.root/"missing.exe")))
        _,document,_=self.artifact(result)
        self.assertEqual(result["status"],"complete")
        self.assertEqual(document["pages"][0]["route"],"blank")

    def test_corrupt_pdf_is_failure_without_success_package(self):
        source=self.root/"corrupt.pdf";source.write_bytes(b"%PDF-1.7\nCORRUPT")
        result=self.job(source)
        self.assertEqual(result["status"],"failed")

    def test_missing_ocr_is_partial_with_visual_evidence(self):
        source,_=self.scan()
        result=self.job(source,Options(tesseract=str(self.root/"missing.exe"),scan_previews=True))
        root,document,md=self.artifact(result)
        self.assertEqual(result["status"],"partial")
        self.assertTrue(any("ocr_failed" in w for p in document["pages"] for w in p["warnings"]))
        assets=[b["asset"] for p in document["pages"] for b in p["blocks"] if b.get("asset")]
        self.assertTrue(assets)
        for asset in assets:self.assertTrue((root/asset).is_file())

    def test_clean_scan_real_ocr_recovers_exact_numbers(self):
        from traslado.ocr import find_tesseract
        exe=find_tesseract()
        if not exe:self.skipTest("Tesseract indisponível: OCR real NOT RUN")
        source,expected=self.scan()
        result=self.job(source,Options(language="eng",tesseract=exe))
        _,document,md=self.artifact(result)
        self.assertFalse(any("ocr_failed" in w for p in document["pages"] for w in p["warnings"]))
        self.assertEqual(document["pages"][0]["route"],"ocr")
        for line in expected:self.assertIn(line,md)

    def test_ocr_job_does_not_mutate_global_tessdata_environment(self):
        source,_=self.scan()
        existed="TESSDATA_PREFIX" in os.environ
        before=os.environ.get("TESSDATA_PREFIX")
        result=self.job(source,Options(tesseract=str(self.root/"missing.exe")))
        self.assertEqual(result["status"],"partial")
        self.assertEqual("TESSDATA_PREFIX" in os.environ,existed)
        self.assertEqual(os.environ.get("TESSDATA_PREFIX"),before)

    def test_reading_order_two_columns_has_each_sentinel_once(self):
        blocks=[Block("text","RIGHT_SECOND",[340,160,540,185]),
                Block("text","LEFT_FIRST",[50,70,250,95]),
                Block("text","TITLE_SENTINEL",[50,15,550,40]),
                Block("text","RIGHT_FIRST",[340,70,540,95]),
                Block("text","LEFT_SECOND",[50,160,250,185])]
        ordered,_=reading_order(blocks,600)
        self.assertEqual([b.text for b in ordered],["TITLE_SENTINEL","LEFT_FIRST","LEFT_SECOND","RIGHT_FIRST","RIGHT_SECOND"])

    def test_pdf_native_table_cells_and_order_have_known_oracle(self):
        import fitz
        doc=fitz.open();page=doc.new_page(width=600,height=800)
        page.insert_text((50,50),"BEFORE_TABLE")
        for x in (50,250,450):page.draw_line((x,100),(x,220))
        for y in (100,160,220):page.draw_line((50,y),(450,y))
        for x,y,text in ((65,130,"Objeto"),(265,130,"Valor"),(65,190,"Servico"),(265,190,"12345")):
            page.insert_text((x,y),text,fontsize=12)
        page.insert_text((50,270),"AFTER_TABLE")
        source=self.root/"table.pdf";doc.save(source);doc.close()
        _,document,md=self.artifact(self.job(source,Options(mode="native")))
        tables=[b["rows"] for p in document["pages"] for b in p["blocks"] if b["kind"]=="table"]
        self.assertTrue(tables)
        self.assertIn(["Servico","12345"],tables[0])
        self.assertLess(md.index("BEFORE_TABLE"),md.index("12345"))
        self.assertLess(md.index("12345"),md.index("AFTER_TABLE"))

    def test_docx_merged_cells_preserve_values_and_spans(self):
        from docx import Document
        doc=Document();doc.add_paragraph("BEFORE_MERGE")
        table=doc.add_table(rows=3,cols=3)
        table.cell(0,0).merge(table.cell(0,1)).text="HORIZONTAL_SENTINEL"
        table.cell(1,0).merge(table.cell(2,0)).text="VERTICAL_SENTINEL"
        table.cell(2,2).text="987654"
        doc.add_paragraph("AFTER_MERGE")
        source=self.root/"merged.docx";doc.save(source)
        _,document,md=self.artifact(self.job(source))
        self.assertEqual(md.count("HORIZONTAL_SENTINEL"),1)
        self.assertEqual(md.count("VERTICAL_SENTINEL"),1)
        self.assertIn('colspan="2"',md)
        self.assertIn('rowspan="2"',md)
        self.assertIn("987654",md)
        self.assertLess(md.index("BEFORE_MERGE"),md.index("987654"))
        self.assertLess(md.index("987654"),md.index("AFTER_MERGE"))

    def test_archive_rejects_traversal_and_symlink_before_parser(self):
        for case in ("traversal","symlink","size"):
            with self.subTest(case=case):
                target=self.root/(case+".docx")
                with zipfile.ZipFile(target,"w",compression=zipfile.ZIP_DEFLATED) as archive:
                    if case=="traversal":archive.writestr("../escape.txt","sentinel")
                    elif case=="symlink":
                        info=zipfile.ZipInfo("link");info.create_system=3;info.external_attr=(stat.S_IFLNK|0o777)<<16
                        archive.writestr(info,"../outside")
                    else:archive.writestr("word/document.xml","X"*2048)
                with self.assertRaises(ValueError):check_archive(target,Options(max_archive_bytes=1024))
                self.assertFalse((self.root.parent/"escape.txt").exists())

    def test_worker_deadline_prevents_late_side_effect(self):
        marker=self.root/"late.txt"
        code="import time,pathlib;time.sleep(5);pathlib.Path("+repr(str(marker))+ ").write_text('LATE')"
        start=time.monotonic()
        with self.assertRaises(TimeoutError):
            run_bounded([sys.executable,"-c",code],self.root,1,self.root/"worker.log")
        self.assertLess(time.monotonic()-start,4)
        self.assertFalse(marker.exists())

    def test_worker_cancellation_prevents_late_side_effect(self):
        marker=self.root/"cancelled.txt"
        code="import time,pathlib;time.sleep(5);pathlib.Path("+repr(str(marker))+ ").write_text('LATE')"
        start=time.monotonic()
        with self.assertRaises(Cancelled):
            run_bounded([sys.executable,"-c",code],self.root,10,self.root/"cancel.log",cancel=lambda:time.monotonic()-start>.25)
        self.assertLess(time.monotonic()-start,4)
        self.assertFalse(marker.exists())


if __name__=="__main__":unittest.main()
