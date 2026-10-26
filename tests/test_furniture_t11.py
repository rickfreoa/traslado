from __future__ import annotations

from dataclasses import asdict
import copy
import json
from pathlib import Path
import random
import sys
import tempfile
import unittest
from unittest.mock import patch

import fitz
from docx import Document as WordDocument

from freoa_v6 import cli, formats, pdf, worker
from freoa_v6.chunks import write_chunks
from freoa_v6.model import Block, Document, Options, Page
from freoa_v6.render import layout
from test_pdf_lines_t11 import baseline_module


ROOT = Path(__file__).resolve().parents[1]
BASELINE = "6a4c4e6ca258dd7fcc0b7a61d0e4dffbe5a3fd70"


class FurnitureT11(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(dir=ROOT / ".t11-logs")
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def source(self):
        path = self.root / "furniture.pdf"
        with fitz.open() as source:
            for i, note in enumerate(("Alpha", "Beta", "Gamma", "Delta"), 1):
                page = source.new_page(width=600, height=800)
                page.insert_text((40, 25), f"Running header - {i}", fontsize=9)
                page.insert_text((300, 55), "Large running title", fontsize=16)
                if i == 1:
                    page.insert_text((40, 65), "Chapter title", fontsize=10)
                page.insert_text((40, 180), "Body prose with measurable font size. " * 2, fontsize=11)
                page.insert_text((40, 740), f"{note} substantive footnote\nDistinct explanation", fontsize=9)
                page.insert_text((290, 790), str(i), fontsize=9)
            source.save(path)
        return path

    def extract(self, path, mode="mark", extractor=None):
        workspace = self.root / f"out-{len(list(self.root.iterdir()))}"
        workspace.mkdir()
        document = Document(path.name, "hash")
        (extractor or formats.extract)(path, workspace, Options(mode="native", furniture=mode), document)
        return document, workspace

    def test_generated_pdf_roles_and_lossless_outputs(self):
        path = self.source()
        marked, workspace = self.extract(path)
        off, _ = self.extract(path, "off")
        for page in marked.pages:
            for block in page.blocks:
                expected = "furniture" if (block.text.startswith("Running header") or block.text.isdecimal()) else "body"
                self.assertEqual(block.role, expected, block.text)
        self.assertEqual(sum(b.role == "furniture" for p in marked.pages for b in p.blocks), 8)
        stripped = copy.deepcopy(marked)
        for page in stripped.pages:
            for block in page.blocks:
                block.role = None
        self.assertEqual(asdict(stripped), asdict(off))
        self.assertEqual(layout(marked).markdown, layout(off).markdown)
        refs = layout(marked).refs
        self.assertEqual([r["role"] for r in refs], [b.role for p in marked.pages for b in p.blocks])
        self.assertEqual([{k: v for k, v in r.items() if k != "role"} for r in refs], layout(off).refs)
        write_chunks(marked, layout(marked).markdown, workspace / "chunks", size_chars=256, overlap_chars=0)
        chunks = [json.loads(line) for line in (workspace / "chunks/chunks.jsonl").read_text(encoding="utf-8").splitlines()]
        self.assertTrue(chunks)
        self.assertTrue(all("role" in b for c in chunks for b in c["blocks"]))

    def test_off_matches_stage5_and_docx_unchanged(self):
        path = self.source()
        off, _ = self.extract(path, "off")
        previous, _ = self.extract(path, "off", baseline_module("pdf", BASELINE).extract)
        self.assertEqual(asdict(off), asdict(previous))
        word = WordDocument()
        word.add_paragraph("Running header 1")
        path = self.root / "source.docx"
        word.save(path)
        marked, _ = self.extract(path)
        off, _ = self.extract(path, "off")
        self.assertEqual(asdict(marked), asdict(off))
        self.assertEqual(layout(marked), layout(off))
        self.assertTrue(all(b.role is None for p in marked.pages for b in p.blocks))

    def test_worker_artifacts_only_add_roles(self):
        from freoa_v6.safety import sha256

        path = self.source()
        outputs = []
        for mode, extractor in (("off", baseline_module("pdf", BASELINE).extract),
                                ("off", pdf.extract), ("mark", pdf.extract)):
            workspace = self.root / f"worker-{len(outputs)}"
            workspace.mkdir()
            request = self.root / "request.json"
            values = asdict(Options(mode="native", furniture=mode, chunks=True))
            request.write_text(json.dumps({"workspace": str(workspace), "snapshot": str(path),
                "source_name": path.name, "sha256": sha256(path), "options": values}), encoding="utf-8")
            with patch.object(sys, "argv", ["worker", str(request)]), patch.object(
                    worker.time, "monotonic", return_value=100), patch.object(pdf, "extract", extractor):
                self.assertEqual(worker.main(), 0)
            outputs.append({p.relative_to(workspace).as_posix(): p.read_bytes()
                            for p in workspace.rglob("*") if p.is_file()})
        self.assertEqual(outputs[0], outputs[1])
        self.assertEqual(outputs[1].keys(), outputs[2].keys())

        def without_roles(value):
            if isinstance(value, dict):
                return {key: without_roles(item) for key, item in value.items() if key != "role"}
            if isinstance(value, list):
                return [without_roles(item) for item in value]
            return value

        for name, payload in outputs[2].items():
            previous = outputs[1][name]
            if name == "document.json":
                data = without_roles(json.loads(payload))
                data["options"]["furniture"] = "off"
                self.assertEqual(data, json.loads(previous))
            elif name == "document-layout.json":
                self.assertEqual(without_roles(json.loads(payload)), json.loads(previous))
            elif name == "chunks/chunks.jsonl":
                self.assertEqual([without_roles(json.loads(line)) for line in payload.splitlines()],
                                 [json.loads(line) for line in previous.splitlines()])
            else:
                self.assertEqual(payload, previous, name)

    def test_options_cli_and_toml(self):
        self.assertEqual(Options().furniture, "mark")
        for invalid in ("omit", "", None, 1):
            with self.subTest(invalid=invalid), self.assertRaises(ValueError):
                Options(furniture=invalid)
        config = self.root / "config.toml"
        config.write_text('[conversion]\nfurniture = "off"\n', encoding="utf-8")
        for flags, expected in (([], "off"), (["--furniture", "mark"], "mark")):
            with patch.object(cli, "discover", return_value=[self.root / "a.pdf"]), patch.object(
                    cli, "execute", return_value={"status": "complete"}) as execute:
                self.assertEqual(cli.main(["convert", "a.pdf", "-o", str(self.root), "--config", str(config), *flags]), 0)
                self.assertEqual(execute.call_args.args[2].furniture, expected)

    def block(self, text, box, size=10, lines=1, engine="pymupdf", kind="text"):
        return Block(kind, text, box, engine=engine, lines=None if lines is None else [
            {"spans": [{"text": text, "size": size}]} for _ in range(lines)])

    def test_candidate_boundaries_repetition_and_excluded_engines(self):
        pages = []
        for i in range(4):
            blocks = [self.block("body" * 40, [0, 100, 100, 700], 11),
                      self.block(f"__ Header {i}!", [0, 0, 100, 80]),
                      self.block("footer", [0, 720, 100, 790]),
                      self.block("XIV", [0, 720, 100, 790]),
                      self.block("12345", [0, 720, 100, 790]),
                      self.block("12", [0, 710, 100, 730]),
                      self.block("7", [0, 720, 100, 790], lines=None),
                      self.block("8", [0, 720, 100, 790], lines=3),
                      self.block("9", [0, 720, 100, 790], size=11.6),
                      self.block("10", [0, 720, 100, 790], size=11.5),
                      self.block("---", [0, 720, 100, 790]),
                      self.block("42", [0, 720, 100, 790], engine="tesseract"),
                      self.block("42", [0, 720, 100, 790], kind="table")]
            # Same text across top/bottom must not combine into three sheets.
            blocks.append(self.block("alternating", [0, 0, 100, 80] if i < 2 else [0, 720, 100, 790]))
            pages.append(Page(i + 1, "native", blocks))
        document = Document("source.pdf", "hash", pages=pages)
        pdf._mark_furniture(document, [800] * 4)
        expected = ["body", "furniture", "furniture", "furniture", "body", "body", "body", "body", "body", "furniture", "body", None, None, "body"]
        for page in pages:
            self.assertEqual([b.role for b in page.blocks], expected)

    def test_no_measurable_body_and_distinct_sheets(self):
        for body in ([], [self.block("body", [0, 100, 100, 700], lines=None)]):
            document = Document("source.pdf", "hash", pages=[Page(1, "native", body + [
                self.block("1", [0, 0, 100, 80]),
                *[self.block("repeated", [0, 0, 100, 80]) for _ in range(4)]])])
            pdf._mark_furniture(document, [800])
            self.assertTrue(all(b.role == "body" for b in document.pages[0].blocks))
        document = Document("source.pdf", "hash", pages=[Page(1, "native", [
            self.block("body", [0, 100, 100, 700]),
            *[self.block("repeated", [0, 0, 100, 80]) for _ in range(4)]])])
        pdf._mark_furniture(document, [800])
        self.assertTrue(all(b.role == "body" for b in document.pages[0].blocks))

    def test_weighted_selection_matches_expanded_median_ranks(self):
        rng = random.Random(42)
        for count in (1, 2, 5, 20, 100, 501):
            weighted = [(rng.randrange(1, 30), rng.randrange(1, 20)) for _ in range(count)]
            expanded = sorted(size for size, weight in weighted for _ in range(weight))
            for rank in ((len(expanded) - 1) // 2, len(expanded) // 2):
                self.assertEqual(pdf._weighted_select(weighted, rank), expanded[rank])
