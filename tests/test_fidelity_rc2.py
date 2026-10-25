"""Source-fidelity regressions reproduced against the frozen rc1 package."""
from __future__ import annotations

from collections import Counter
import html
import json
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import pymupdf

from freoa_v6.model import Block, Document as OutputDocument, Options
from freoa_v6.pdf import reading_order, extract as extract_pdf
from freoa_v6.render import render
from freoa_v6.runner import execute


def xml_paragraph(text):
    p = OxmlElement("w:p")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = text
    r.append(t)
    p.append(r)
    return p


class SourceFidelityRc2(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="freoa-fidelity-")
        self.root = Path(self.temp.name)

    def tearDown(self):
        self.temp.cleanup()

    def convert(self, path):
        result = execute(path, self.root / "results", Options(mode="native"))
        self.assertIn(result["status"], {"complete", "partial"}, result)
        package = Path(result["output"])
        data = json.loads((package / "document.json").read_text(encoding="utf-8"))
        markdown = (package / "document.md").read_text(encoding="utf-8")
        return result, data, markdown

    def save_docx(self, doc, name="record.docx"):
        path = self.root / name
        doc.save(path)
        return path

    def warnings(self, data):
        return [w for page in data["pages"] for w in page["warnings"]]

    def test_docx_table_break_tab_and_carriage_return_preserve_separation(self):
        doc = Document()
        table = doc.add_table(rows=1, cols=3)
        first = table.cell(0, 0).paragraphs[0]
        first.add_run("123")
        first.add_run().add_break()
        first.add_run("456")
        second = table.cell(0, 1).paragraphs[0]
        second.add_run("DEBT")
        second.add_run().add_tab()
        second.add_run("100")
        third = table.cell(0, 2).paragraphs[0]
        third.add_run("789")
        third.add_run()._r.append(OxmlElement("w:cr"))
        third.add_run("012")
        path = self.save_docx(doc)
        # Independent source oracle before conversion.
        source = Document(path)
        self.assertEqual(source.tables[0].cell(0, 0).text, "123\n456")
        self.assertEqual(source.tables[0].cell(0, 1).text, "DEBT\t100")
        _, data, markdown = self.convert(path)
        raw = "\n".join(block["text"] for page in data["pages"] for block in page["blocks"])
        for text in (raw, markdown):
            self.assertIn("123<br>456", text)
            self.assertIn("DEBT\t100", html.unescape(text))
            self.assertIn("789<br>012", text)
            self.assertNotIn("123456", text)
            self.assertNotIn("DEBT100", text)
            self.assertNotIn("789012", text)

    def test_nested_docx_table_text_occurs_once(self):
        doc = Document()
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        cell.text = "OUTER_BEFORE"
        nested = cell.add_table(rows=2, cols=2)
        values = ["INNER_HEADER_A", "INNER_HEADER_B", "INNER_VALUE_A", "987654"]
        for index, value in enumerate(values):
            nested.cell(index // 2, index % 2).text = value
        cell.add_paragraph("OUTER_AFTER")
        _, data, markdown = self.convert(self.save_docx(doc))
        for token in ["OUTER_BEFORE", *values, "OUTER_AFTER"]:
            self.assertEqual(markdown.count(token), 1, token)
        self.assertIn("docx_nested_table_structure_flattened", self.warnings(data))

    def test_interrupted_vertical_merge_preserves_source_cell(self):
        doc = Document()
        table = doc.add_table(rows=3, cols=2)
        table.cell(0, 0).text = "ANCHOR_TOP"
        restart = OxmlElement("w:vMerge")
        restart.set(qn("w:val"), "restart")
        table._tbl.tr_lst[0].tc_lst[0].get_or_add_tcPr().append(restart)
        table.cell(1, 0).merge(table.cell(1, 1)).text = "INTERRUPTING_WIDE_CELL"
        last = table._tbl.tr_lst[2].tc_lst[0]
        last.clear_content()
        last.append(xml_paragraph("BOTTOM_98765"))
        last.get_or_add_tcPr().append(OxmlElement("w:vMerge"))
        _, data, markdown = self.convert(self.save_docx(doc))
        self.assertEqual(markdown.count("BOTTOM_98765"), 1)
        self.assertIn("BOTTOM_98765", json.dumps(data))
        self.assertNotIn('rowspan="2"', markdown)
        self.assertIn("docx_vmerge_inconsistent_content_preserved", self.warnings(data))

    def test_adjacent_vertical_merge_preserves_nonempty_continuation(self):
        doc = Document()
        table = doc.add_table(rows=2, cols=1)
        table.cell(0, 0).text = "TOP_SOURCE"
        table.cell(1, 0).text = "CONTINUATION_98765"
        restart = OxmlElement("w:vMerge")
        restart.set(qn("w:val"), "restart")
        table._tbl.tr_lst[0].tc_lst[0].get_or_add_tcPr().append(restart)
        table._tbl.tr_lst[1].tc_lst[0].get_or_add_tcPr().append(OxmlElement("w:vMerge"))
        _, data, markdown = self.convert(self.save_docx(doc))
        self.assertEqual(markdown.count("CONTINUATION_98765"), 1)
        self.assertIn("TOP_SOURCE<br>CONTINUATION_98765", markdown)
        self.assertIn('rowspan="2"', markdown)
        self.assertIn("docx_vmerge_continuation_text_preserved", self.warnings(data))

    def test_plain_row_ends_vertical_merge_chain(self):
        doc = Document()
        table = doc.add_table(rows=3, cols=1)
        for index, text in enumerate(("TOP_SOURCE", "UNMERGED_MIDDLE", "ORPHAN_BOTTOM")):
            table.cell(index, 0).text = text
        restart = OxmlElement("w:vMerge")
        restart.set(qn("w:val"), "restart")
        table._tbl.tr_lst[0].tc_lst[0].get_or_add_tcPr().append(restart)
        table._tbl.tr_lst[2].tc_lst[0].get_or_add_tcPr().append(OxmlElement("w:vMerge"))
        _, data, markdown = self.convert(self.save_docx(doc))
        self.assertNotIn('rowspan="2"', markdown)
        self.assertEqual(markdown.count("ORPHAN_BOTTOM"), 1)
        self.assertIn("docx_vmerge_inconsistent_content_preserved", self.warnings(data))

    def test_pdf_narrow_centered_digit_occurs_once_in_effective_outputs(self):
        doc = pymupdf.open()
        page = doc.new_page(width=600, height=800)
        for x, y, text in [(45,75,"LEFT_FIRST_PAYABLE_SENTINEL"), (45,170,"LEFT_SECOND_PAYABLE_SENTINEL"),
                           (340,75,"RIGHT_FIRST_PAYABLE_SENTINEL"), (340,170,"RIGHT_SECOND_PAYABLE_SENTINEL")]:
            page.insert_text((x, y), text, fontsize=9)
        page.insert_text((299, 125), "9", fontsize=4)
        path = self.root / "center.pdf"
        doc.save(path)
        doc.close()
        _, data, markdown = self.convert(path)
        self.assertEqual(data["pages"][0]["native_text"].split().count("9"), 1)
        self.assertEqual(sum(b["text"] == "9" for b in data["pages"][0]["blocks"]), 1)
        self.assertEqual(sum(line == "9" for line in markdown.splitlines()), 1)

    def test_reading_order_preserves_each_input_record_including_equal_values(self):
        blocks = [Block("text", "LEFT_EQUAL", [40,70,240,95]),
                  Block("text", "LEFT_EQUAL", [40,70,240,95]),
                  Block("text", "RIGHT_FIRST", [340,70,540,95]),
                  Block("text", "RIGHT_SECOND", [340,170,540,195]),
                  Block("text", "9", [299,100,301,110]),
                  Block("text", "NO_COORDINATES")]
        ordered, _ = reading_order(blocks, 600)
        self.assertEqual(Counter(id(b) for b in ordered), Counter(id(b) for b in blocks))
        self.assertEqual(sum(b.text == "9" for b in ordered), 1)

    def test_nested_body_sdt_preserves_paragraph_and_table_order(self):
        doc = Document()
        doc.add_paragraph("BODY_BEFORE")
        table = doc.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "SDT_TABLE_54321"
        sdt, content = OxmlElement("w:sdt"), OxmlElement("w:sdtContent")
        content.append(xml_paragraph("SDT_SENTINEL_98765"))
        content.append(table._tbl)
        nested, nested_content = OxmlElement("w:sdt"), OxmlElement("w:sdtContent")
        nested_content.append(xml_paragraph("NESTED_SDT_SENTINEL"))
        nested.append(nested_content)
        content.append(nested)
        sdt.append(content)
        doc.element.body.insert(1, sdt)
        doc.add_paragraph("BODY_AFTER")
        _, data, markdown = self.convert(self.save_docx(doc))
        sentinels = ["BODY_BEFORE", "SDT_SENTINEL_98765", "SDT_TABLE_54321", "NESTED_SDT_SENTINEL", "BODY_AFTER"]
        positions = [markdown.index(token) for token in sentinels]
        self.assertEqual(positions, sorted(positions))
        for token in sentinels:
            self.assertEqual(markdown.count(token), 1)
            self.assertIn(token, json.dumps(data))

    def test_tracked_insertions_preserved_consistently_with_unresolved_warning(self):
        doc = Document()
        paragraph = doc.add_paragraph("BODY_PREFIX ")
        table_paragraph = doc.add_table(rows=1, cols=1).cell(0,0).paragraphs[0]
        table_paragraph.add_run("CELL_PREFIX ")
        for p, text in ((paragraph, "BODY_INSERTED_12345"), (table_paragraph, "CELL_INSERTED_67890")):
            ins = OxmlElement("w:ins")
            ins.set(qn("w:id"), "1")
            run = OxmlElement("w:r")
            element = OxmlElement("w:t")
            element.text = text
            run.append(element)
            ins.append(run)
            p._p.append(ins)
            p.add_run(" SUFFIX")
        _, data, markdown = self.convert(self.save_docx(doc))
        for token in ("BODY_INSERTED_12345", "CELL_INSERTED_67890"):
            self.assertEqual(markdown.count(token), 1)
        self.assertIn("docx_tracked_changes_not_resolved", self.warnings(data))

    def test_pdf_form_and_annotations_require_explicit_review(self):
        doc = pymupdf.open()
        page = doc.new_page(width=600, height=800)
        page.insert_text((50,75), "NATIVE_LABEL")
        widget = pymupdf.Widget()
        widget.field_name = "PaymentAmount"
        widget.field_type = pymupdf.PDF_WIDGET_TYPE_TEXT
        widget.field_value = "FORM_VALUE_456789"
        widget.rect = pymupdf.Rect(50,120,400,180)
        page.add_widget(widget)
        page.add_text_annot((50,230), "ANNOT_SENTINEL_12345")
        path = self.root / "form.pdf"
        doc.save(path)
        doc.close()
        result, data, markdown = self.convert(path)
        self.assertEqual(result["status"], "partial")
        self.assertIn("pdf_form_field_semantics_unverified", self.warnings(data))
        self.assertIn("pdf_annotations_not_extracted", self.warnings(data))
        self.assertIn("NATIVE_LABEL", data["pages"][0]["native_text"])
        self.assertIn("FORM_VALUE_456789", markdown)

    def controlled_table_extract(self, table_cells):
        # Real PDF text/geometry, controlled finder output: this isolates handling
        # of an upstream omission without claiming the real finder caused one.
        doc = pymupdf.open()
        page = doc.new_page(width=600, height=800)
        for x in (50,250,450):
            page.draw_line((x,110), (x,220))
        for y in (110,165,220):
            page.draw_line((50,y), (450,y))
        page.insert_text((65,135), "PAYABLE_SENTINEL", fontsize=11)
        page.insert_text((265,135), "12345", fontsize=11)
        path = self.root / "table-loss.pdf"
        doc.save(path)
        doc.close()
        tables=[]
        for cells in table_cells:
            tables.append(SimpleNamespace(bbox=(50,110,450,220),
                          extract=lambda cells=cells: cells,
                          rows=[SimpleNamespace(cells=[(50,110,250,165),(250,110,450,165)])]))
        document=OutputDocument(path.name, "synthetic-oracle")
        workspace=self.root/"controlled"
        workspace.mkdir()
        with patch.object(pymupdf.Page, "find_tables", return_value=SimpleNamespace(tables=tables)):
            extract_pdf(path,workspace,Options(mode="native"),document)
        return document, render(document)

    def test_upstream_table_omission_retains_native_source_text(self):
        document, markdown=self.controlled_table_extract([[["", "12345"]]])
        self.assertIn("PAYABLE_SENTINEL",markdown)
        self.assertTrue(any(b.kind=="text" and "PAYABLE_SENTINEL" in b.text for b in document.pages[0].blocks))
        self.assertIn("pdf_table_cell_text_mismatch_preserved",document.pages[0].warnings)

    def test_second_table_cannot_delete_previously_extracted_table(self):
        document, markdown=self.controlled_table_extract([[["PAYABLE_SENTINEL","12345"]], [["SECOND_TABLE_ONLY","67890"]]])
        tables=[b for b in document.pages[0].blocks if b.kind=="table"]
        self.assertEqual(len(tables),2)
        self.assertIn("PAYABLE_SENTINEL",markdown)
        self.assertIn("SECOND_TABLE_ONLY",markdown)
        self.assertIn("pdf_overlapping_table_regions_need_review",document.pages[0].warnings)


if __name__ == "__main__":
    unittest.main()
