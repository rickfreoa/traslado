# SPDX-License-Identifier: AGPL-3.0-or-later
"""Regressions for all ten findings in the independent E2/E3 review."""
import io
import json
import os
from pathlib import Path
import signal
from types import SimpleNamespace
from unittest.mock import patch
import zipfile
import zlib

import pytest

from traslado import detect, formats, platform, runner, safety
from traslado.errors import ConversionError, classify, partial_error
from traslado.model import Document, Options


@pytest.mark.parametrize('suffix,payload', [
    ('.txt', 'BM&FBOVESPA: relatorio de mercado\nLinha 2\n'),
    ('.md', 'BM&FBOVESPA: relatorio de mercado\nLinha 2\n'),
    ('.csv', 'BM,valor\nacao,123\n'),
    ('.md', '# Brief\nDetectar PDF (`%PDF-` nos primeiros 1024 bytes).\n'),
    ('.md', '# Brief\nExemplo %PDF-1.7 em prosa.\n'),
])
def test_text_false_positives(tmp_path, suffix, payload):
    source = tmp_path / ('record' + suffix)
    source.write_text(payload, encoding='utf-8')
    assert detect.sniff(source) is None
    result = runner.execute(source, tmp_path / 'out', Options(mode='native'))
    assert result['status'] == 'complete', result
    assert 'error_code' not in result


@pytest.mark.parametrize('suffix,offset,expected', [('.md', 1, None), ('.txt', 0, '.pdf'),
    ('.bin', 1024, '.pdf'), ('.bin', 1025, None)])
def test_pdf_offset_policy(tmp_path, suffix, offset, expected):
    source = tmp_path / ('record' + suffix)
    source.write_bytes(b'x' * offset + b'%PDF-1.7')
    assert detect.sniff(source) == expected


def rewrite_member(path, member, payload):
    with zipfile.ZipFile(path) as archive:
        members = {n: archive.read(n) for n in archive.namelist()}
    members[member] = payload
    with zipfile.ZipFile(path, 'w') as archive:
        for n, data in members.items():
            archive.writestr(n, data)


@pytest.mark.parametrize('kind', ['docx', 'pptx', 'xlsx', 'workbook', 'bmp'])
def test_real_malformed_parsers(tmp_path, kind):
    from docx import Document as Docx
    from pptx import Presentation
    from openpyxl import Workbook
    source = tmp_path / ('record.' + ('xlsx' if kind == 'workbook' else kind))
    expected = 'malformed'
    if kind == 'docx':
        Docx().save(source)
        rewrite_member(source, 'word/document.xml', b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"/>')
    elif kind == 'pptx':
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[6])
        prs.save(source)
        rewrite_member(source, 'ppt/slides/slide1.xml', b'<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"/>')
    elif kind == 'xlsx':
        wb = Workbook()
        wb.active['A1'] = 'sentinel'
        wb.save(source)
        rewrite_member(source, 'xl/worksheets/sheet1.xml', b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row r="1"><c r="A1" t="s"><v>999</v></c></row></sheetData></worksheet>')
    elif kind == 'workbook':
        with zipfile.ZipFile(source, 'w') as archive:
            archive.writestr('[Content_Types].xml', '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>')
            archive.writestr('xl/workbook.xml', '<workbook/>')
        expected = 'missing_part'
    else:
        source.write_bytes(b'BM' + bytes(12) + (40).to_bytes(4, 'little') + bytes(4))
    result = runner.execute(source, tmp_path / 'out', Options(mode='native'))
    assert result.get('error_code') == expected, result


@pytest.mark.parametrize('exc', [IndexError('index'), KeyError('key'), AttributeError('attr'), TypeError('type')])
def test_only_parser_calls_are_mapped(exc):
    with pytest.raises(ConversionError) as error:
        formats._parser_call(lambda: (_ for _ in ()).throw(exc))
    assert error.value.code == 'malformed'
    assert classify(exc).code == 'internal'


def test_pdf_localized_runtime(tmp_path):
    import fitz
    from traslado.pdf import extract
    source = tmp_path / 'record.pdf'
    with fitz.open() as pdf:
        page = pdf.new_page()
        page.insert_text((40, 40), 'Native sentinel text 12345')
        for y in (80, 100, 120):
            page.draw_line((40, y), (200, y))
        pdf.save(source)
    doc = Document(source.name, 'hash')
    (tmp_path / 'assets').mkdir()
    with patch.object(fitz.Page, 'find_tables', side_effect=RuntimeError('MuPDF page failure')):
        extract(source, tmp_path, Options(mode='native'), doc)
    assert partial_error(doc).code == 'malformed'


def test_partial_fidelity_has_no_error(tmp_path):
    source = tmp_path / 'record.html'
    source.write_text('<html><p>sentinel</p></html>')
    result = runner.execute(source, tmp_path / 'out')
    assert result['status'] == 'partial'
    assert 'error_code' not in result
    assert partial_error(Document('x', 'hash')) is None


def candidate(root, release='v1', gates=None):
    (root / 'releases' / release).mkdir(parents=True, exist_ok=True)
    platform.write_candidate(root, {'release_id': release, 'gates': gates if gates is not None else {'doctor': True, 'fixture': True}})


def test_where_absent_and_invalid(tmp_path):
    assert platform.where(tmp_path, platform='linux')['errors'] == []
    for payload in ('[]', '{}', 'broken'):
        (tmp_path / 'candidate.json').write_text(payload)
        assert platform.where(tmp_path, platform='linux')['errors']


@pytest.mark.parametrize('gates', [None, [], {'doctor': 1, 'fixture': True}] + [
    {'doctor': True, 'fixture': True, 'revocation': value} for value in [False, 'erro', [], {}]])
def test_promotion_invalid_gates(tmp_path, gates):
    candidate(tmp_path)
    data = platform.read_candidate(tmp_path)
    data['gates'] = gates
    platform.write_candidate(tmp_path, data)
    with pytest.raises(ValueError, match='Gates'):
        platform.promote(tmp_path)
    assert not (tmp_path / 'current.json').exists()
    assert not (tmp_path / 'state/promote.lock').exists()


def test_promotion_corrupt_current_repromotion_and_lock(tmp_path, caplog):
    candidate(tmp_path)
    current = tmp_path / 'current.json'
    current.write_text('corrupted')
    with pytest.raises(ValueError, match='current.json'):
        platform.promote(tmp_path)
    assert current.read_text() == 'corrupted'
    current.unlink()
    platform.promote(tmp_path)
    candidate(tmp_path, 'v2')
    platform.promote(tmp_path)
    candidate(tmp_path, 'v2')
    assert platform.promote(tmp_path)['previous_release_id'] == 'v1'
    lock = tmp_path / 'state/promote.lock'
    candidate(tmp_path, 'v3')
    with platform._promotion_lock(tmp_path):
        for action in (platform.promote, platform.rollback):
            with pytest.raises(ValueError, match='promote.lock'):
                action(tmp_path)
    lock.write_text('stale')
    os.utime(lock, (0, 0))
    assert platform.promote(tmp_path)['release_id'] == 'v3'
    assert 'obsoleto' in caplog.text
    assert not lock.exists()


def test_injected_home_shortcuts(tmp_path, monkeypatch):
    foreign = tmp_path / 'foreign'
    home = tmp_path / 'home'
    monkeypatch.setenv('APPDATA', str(foreign))
    with patch.object(platform, '_desktop', wraps=platform._desktop) as desk, patch.object(platform.subprocess, 'run'):
        paths = platform.create_shortcut('python', platform='win32', home=home)
        assert desk.call_args.args == (home, 'win32', True)
        assert all(path.is_relative_to(home) for path in paths)
    assert not foreign.exists()
    assert platform.remove_shortcuts(['Traslado'], platform='win32', home=home) == []


def test_desktop_escaping_names_and_duplicates(tmp_path):
    value = 'a\\b"c$d`e%f'
    encoded = platform._desktop_exec(value)
    # Undo general string escaping, then Exec quoting (independent decoder).
    import re
    decoded = encoded.replace('\\\\', '\\')
    decoded = re.sub(r'\\([\\"$`])', r'\1', decoded).replace('%%', '%')
    assert decoded == value
    paths = platform.create_shortcut('python', name='Traslado Revisão', platform='linux', home=tmp_path)
    assert paths[0].name == 'traslado-revisao.desktop'
    assert platform.remove_shortcuts(['Traslado Revisão', 'Traslado Revisão'], platform='linux', home=tmp_path) == paths


@pytest.mark.parametrize('member,limit', [('mimetype', 256), ('[Content_Types].xml', 65536)])
def test_sniff_zip_physical_limit(tmp_path, member, limit):
    source = tmp_path / 'record.bin'
    with zipfile.ZipFile(source, 'w') as archive:
        archive.writestr(member, b'x')
    reads = []
    class Stream(io.BytesIO):
        def read(self, size=-1):
            reads.append(size)
            return super().read(size)
    with patch.object(zipfile.ZipFile, 'open', return_value=Stream(b'x' * (limit + 100))):
        assert detect.sniff(source) is None
    assert reads == [limit + 1]
    for exc in (zlib.error('deflate'), EOFError(), RuntimeError(), NotImplementedError()):
        with patch.object(zipfile.ZipFile, 'open', side_effect=exc):
            assert detect.sniff(source) is None


def test_proc_errors_are_per_pid(tmp_path):
    root = tmp_path / 'proc'
    root.mkdir()
    for name in ('1', '2', '3', '4', '5'):
        (root / name).mkdir()
    real_path = Path
    def read(path, *args, **kwargs):
        pid = path.parent.name
        if pid == '1':
            raise PermissionError('hidepid')
        if pid == '2':
            return 'invalid'
        if pid == '3':
            return '1 (name) R bad invalid'
        return '5 (name) R 1 99' if path.name == 'stat' else '100 4'
    with patch.object(safety.sys, 'platform', 'linux'), patch.object(safety, 'Path', side_effect=lambda p: root if p == '/proc' else real_path(p)), patch.object(Path, 'read_text', read), patch.object(os, 'sysconf', create=True, return_value=4096):
        assert safety.group_processes(99) == [(4, 16384), (5, 16384)]
    with patch.object(safety.sys, 'platform', 'darwin'), patch.object(safety.subprocess, 'check_output', return_value='5 99 4 S extra\n'):
        assert safety.group_processes(99) == [(5, 4096)]


@pytest.mark.parametrize('code,log,measurements,expected', [
    (-11, b'', {}, 'internal'), (-6, b'bad_alloc', {}, 'resource_limit'),
    (-9, b'out of memory', {}, 'resource_limit'), (-11, b'cannot allocate memory', {}, 'resource_limit'),
    (-11, b'', {'rlimit_applied': True, 'peak_bytes': 95, 'limit_bytes': 100}, 'resource_limit'),
    (-6, b'', {'rlimit_applied': True, 'peak_bytes': 10, 'limit_bytes': 100}, 'internal'),
    (-9, b'', {'detected_after_s': .1}, 'resource_limit')])
def test_native_memory_failure(tmp_path, code, log, measurements, expected):
    source = tmp_path / 'record.txt'
    source.write_text('sentinel')
    def run(argv, cwd, timeout, log_path, *args, **kwargs):
        log_path.write_bytes(log)
        kwargs['measurements'].update(measurements)
        return code
    with patch.object(runner, 'run_bounded', side_effect=run):
        assert runner.execute(source, tmp_path / 'out')['error_code'] == expected


def test_ocr_passes_memory_budget(tmp_path):
    from traslado import ocr
    options = Options(max_memory_mb=384)
    seen = []
    def run(argv, cwd, timeout, log, **kwargs):
        seen.append(kwargs['max_memory_bytes'])
        log.write_text('por\neng\n')
        return 1  # both operations stop after the bounded launch
    with patch.object(ocr, 'run_bounded', side_effect=run):
        with pytest.raises(ConversionError):
            ocr._validate_languages('tesseract', None, options, tmp_path)
    with patch.object(ocr, 'find_tesseract', return_value='tesseract'), patch.object(ocr, 'managed_tessdata_for', return_value=None), patch.object(ocr, '_validate_languages', return_value=['por', 'eng']), patch.object(ocr, 'run_bounded', side_effect=run):
        with pytest.raises(ConversionError):
            ocr._recognize(tmp_path / 'image.png', tmp_path, options)
    assert seen == [384 * 1024 * 1024] * 2
