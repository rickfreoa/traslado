# SPDX-License-Identifier: AGPL-3.0-or-later
# Copyright (C) 2026 Ricardo Peres Freoa
"""A5a: notes are content, resolved through document relationships."""
from __future__ import annotations

from pathlib import Path
import re
import shutil
import unittest
import uuid

from docx import Document
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from traslado import formats, registry
from traslado.model import Block, Document as OutputDocument, Options, Page
from traslado.render import render


def element(tag, text=None, **attrs):
    node = OxmlElement('w:' + tag)
    for key, value in attrs.items():
        node.set(qn('w:' + key), str(value))
    if text is not None:
        node.text = text
    return node


def paragraph(text):
    p = element('p')
    r = element('r')
    r.append(element('t', text))
    p.append(r)
    return p


def reference(p, family, number, custom=False, on_run=False):
    r = element('r')
    ref = element(family + 'Reference', id=number)
    if custom:
        (r if on_run else ref).set(qn('w:customMarkFollows'), '1')
    r.append(ref)
    p.append(r)


def notes(doc, family, definitions):
    root = element(family + 's')
    for number, content, kind in definitions:
        note = element(family, id=number)
        if kind:
            note.set(qn('w:type'), kind)
        for node in content:
            note.append(node)
        root.append(note)
    # Deliberately nonstandard names: only relationship type is authoritative.
    from lxml import etree
    part = Part(PackURI('/word/other-' + family + '.xml'),
                'application/vnd.openxmlformats-officedocument.wordprocessingml.' + family + 's+xml',
                etree.tostring(root), doc.part.package)
    doc.part.relate_to(part, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/' + family + 's')


class DocxNotesT11(unittest.TestCase):
    def setUp(self):
        # Ordinary workspace directory avoids Windows Python 3.14 temp ACLs.
        self.root = Path(__file__).resolve().parents[1] / '.t11-logs' / ('notes-' + uuid.uuid4().hex)
        self.root.mkdir(parents=True)

    def tearDown(self):
        shutil.rmtree(self.root)

    def extract(self, doc):
        path = self.root / 'sample.docx'
        doc.save(path)
        output = OutputDocument(path.name, 'hash')
        formats._extract_docx(path, self.root, Options(images=False), output)
        return output, output.pages[0]

    def test_families_order_separators_and_paragraphs(self):
        doc = Document()
        p = doc.add_paragraph('Corpo ')._p
        for family, number in [('endnote', 2), ('footnote', 3), ('footnote', 1), ('footnote', 3)]:
            reference(p, family, number)
        first = paragraph('Primeiro')
        first.insert(0, element('r'))
        first[0].append(element('footnoteRef'))
        # Even text in the numbering run must not leak into the definition.
        first[0].append(element('t', 'NUMERO'))
        notes(doc, 'footnote', [(1, [paragraph('Um')], None), (3, [first, paragraph('Segundo')], None)] +
              [(n, [paragraph('SEPARADOR')], kind) for n, kind in
               [(-1, 'separator'), (0, 'continuationSeparator'), (9, 'continuationNotice')]])
        notes(doc, 'endnote', [(2, [paragraph('Fim')], None)])
        output, page = self.extract(doc)
        self.assertEqual(page.blocks[0].text, 'Corpo [^e2][^3][^1][^3]')
        self.assertEqual([b.text for b in page.blocks[1:]],
                         ['[^3]: Primeiro\n\n    Segundo', '[^1]: Um', '[^e2]: Fim'])
        for block, family, number in zip(page.blocks[1:], ['footnote', 'footnote', 'endnote'], ['3', '1', '2']):
            self.assertEqual(block.kind, 'text')
            self.assertEqual(block.engine, 'python-docx')
            self.assertEqual(block.details, {'note_family': family, 'note_id': number})
        self.assertEqual(page.warnings, ['docx_formatting_partial'])
        self.assertIn('[^3]: Primeiro\n\n    Segundo', render(output))
        self.assertEqual(len(output.pages), 1)

    def test_table_sdt_and_first_call_order(self):
        doc = Document()
        cell = doc.add_table(rows=1, cols=1).cell(0, 0)
        cell.text = 'Célula'
        reference(cell.paragraphs[0]._p, 'footnote', 8)
        sdt = element('sdt')
        content = element('sdtContent')
        p = paragraph('Controle')
        reference(p, 'footnote', 4)
        content.append(p)
        sdt.append(content)
        doc.element.body.insert(-1, sdt)
        notes(doc, 'footnote', [(4, [paragraph('Quatro')], None), (8, [paragraph('Oito')], None)])
        _, page = self.extract(doc)
        self.assertIn('Célula[^8]', page.blocks[0].text)
        self.assertEqual(page.blocks[1].text, 'Controle[^4]')
        self.assertEqual([b.text for b in page.blocks[2:]], ['[^8]: Oito', '[^4]: Quatro'])

    def test_unreferenced_and_missing_definitions(self):
        doc = Document()
        p = doc.add_paragraph('Texto')._p
        reference(p, 'footnote', 7)
        reference(p, 'endnote', 7)
        reference(p, 'footnote', 2)
        notes(doc, 'footnote', [(5, [paragraph('Órfã')], None), (2, [paragraph('Dois')], None),
                                (6, [], None)])
        notes(doc, 'endnote', [(5, [paragraph('Órfã final')], None)])
        _, page = self.extract(doc)
        self.assertEqual(page.blocks[0].text, 'Texto[nota 7 não encontrada][nota de fim 7 não encontrada][^2]')
        self.assertEqual([b.text for b in page.blocks[1:]],
                         ['[^2]: Dois', '[^5]: Órfã', '[^6]: ', '[^e5]: Órfã final'])
        self.assertEqual(set(page.warnings), {'docx_note_definition_unreferenced',
                                            'docx_note_reference_unresolved', 'docx_formatting_partial'})

    def test_duplicate_missing_and_unsafe_ids_lose_no_definition(self):
        doc = Document()
        p = doc.add_paragraph('Texto')._p
        reference(p, 'footnote', 'abc')
        reference(p, 'footnote', 'a b')
        notes(doc, 'footnote', [('abc', [paragraph('N')], None), ('abc', [paragraph('dup')], None),
                                ('a b', [paragraph('Espaço')], None)])
        root = element('footnote')  # definition without w:id cannot be called, but its text survives
        root.append(paragraph('Sem id'))
        part = next(r.target_part for r in doc.part.rels.values() if r.reltype.endswith('/footnotes'))
        from lxml import etree
        tree = etree.fromstring(part.blob)
        tree.append(root)
        part._blob = etree.tostring(tree)
        _, page = self.extract(doc)
        self.assertEqual(page.blocks[0].text, 'Texto[^abc][^a_b]')
        texts = [b.text for b in page.blocks[1:]]
        self.assertEqual(texts, ['[^abc]: N', '[^a_b]: Espaço', '[^abc-extra1]: dup', '[^sem-id-extra2]: Sem id'])
        self.assertEqual(page.blocks[2].details['note_id'], 'a b')
        self.assertIn('docx_note_duplicate_or_missing_id', page.warnings)

    def test_custom_mark_after_following_run(self):
        for family in ('footnote', 'endnote'):
            for on_run in (False, True):
                with self.subTest(family=family, on_run=on_run):
                    doc = Document()
                    p = doc.add_paragraph('Antes')._p
                    reference(p, family, 1, custom=True, on_run=on_run)
                    p.append(paragraph('*')[0])
                    p.append(paragraph(' depois')[0])
                    notes(doc, family, [(1, [paragraph('Nota')], None)])
                    _, page = self.extract(doc)
                    marker = '[^1]' if family == 'footnote' else '[^e1]'
                    self.assertEqual(page.blocks[0].text, 'Antes*' + marker + ' depois')

    def test_reference_without_any_note_part(self):
        doc = Document()
        p = doc.add_paragraph('Texto')._p
        reference(p, 'footnote', 12)
        reference(p, 'endnote', 13, custom=True)
        p.append(paragraph('†')[0])
        _, page = self.extract(doc)
        self.assertEqual(page.blocks[0].text,
                         'Texto[nota 12 não encontrada]†[nota de fim 13 não encontrada]')
        self.assertEqual(page.warnings, ['docx_note_reference_unresolved', 'docx_formatting_partial'])

    def test_all_xml_calls_and_definitions_survive_in_markdown(self):
        doc = Document()
        for family in ('endnote', 'footnote'):
            for number in [3, 1, 3, 2]:
                p = doc.add_paragraph('Chamada ')._p
                reference(p, family, number)
            notes(doc, family, [(n, [paragraph(f'{family} {n}')], None) for n in [1, 2, 3, 4]])
        output, page = self.extract(doc)
        markdown = render(output)
        for family, prefix in [('footnote', ''), ('endnote', 'e')]:
            count = sum(1 for _ in doc.element.body.iter(qn('w:' + family + 'Reference')))
            calls = re.findall(r'\[\^' + prefix + r'\d+\](?!:)', markdown)
            self.assertEqual(len(calls), count)
            self.assertEqual([b.details['note_id'] for b in page.blocks if b.details.get('note_family') == family],
                             ['3', '1', '2', '4'])
            for number in [1, 2, 3, 4]:
                self.assertEqual(markdown.count(f'[^{prefix}{number}]: {family} {number}'), 1)

    def test_complex_note_xml_order_and_no_numbering_run(self):
        doc = Document()
        reference(doc.add_paragraph('Texto')._p, 'footnote', 1)
        table = element('tbl')
        row, cell = element('tr'), element('tc')
        cell.append(paragraph('Tabela'))
        row.append(cell)
        table.append(row)
        p = paragraph('Antes da caixa')
        r, drawing, box = element('r'), element('drawing'), element('txbxContent')
        box.append(paragraph('Caixa'))
        drawing.append(box)
        r.append(drawing)
        p.append(r)
        p.append(paragraph('Depois da caixa')[0])
        notes(doc, 'footnote', [(1, [paragraph('Início'), table, p, paragraph('Fim')], None)])
        _, page = self.extract(doc)
        self.assertEqual(page.blocks[-1].text,
                         '[^1]: Início\n\n    Tabela\n\n    Antes da caixa\n\n    Caixa\n\n    Depois da caixa\n\n    Fim')
        self.assertIn('docx_note_complex_content_flattened', page.warnings)

    def test_headers_footers_and_comments_only_when_present(self):
        for text, comments in [('', False), ('  ', False), ('Cabeçalho', False), ('', True)]:
            with self.subTest(text=text, comments=comments):
                doc = Document()
                doc.add_paragraph('Corpo')
                doc.sections[0].header.paragraphs[0].text = text
                doc.sections[0].footer.paragraphs[0].text = text
                if comments:
                    doc.add_comment(doc.paragraphs[0].runs[0], text='Comentário', author='Teste')
                _, page = self.extract(doc)
                expected = {'docx_formatting_partial'}
                if text.strip():
                    expected.add('docx_headers_footers_not_extracted')
                if comments:
                    expected.add('docx_comments_not_extracted')
                self.assertEqual(set(page.warnings), expected)

    def test_empty_comments_part_does_not_warn(self):
        from lxml import etree
        doc = Document()
        doc.add_paragraph('Corpo')
        root = element('comments')
        part = Part(PackURI('/word/other-comments.xml'),
                    'application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml',
                    etree.tostring(root), doc.part.package)
        doc.part.relate_to(part, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments')
        _, page = self.extract(doc)
        self.assertEqual(page.warnings, ['docx_formatting_partial'])

    def test_plain_docx_unchanged_except_aggregate_warning(self):
        doc = Document()
        doc.add_heading('Título', level=1)
        doc.add_paragraph('Texto\tcom quebra\ncontinua.')
        _, page = self.extract(doc)
        expected = Page(1, 'native', [Block('text', '# Título', engine='python-docx'),
                                    Block('text', 'Texto\tcom quebra\ncontinua.', engine='python-docx')],
                        ['docx_formatting_partial'])
        self.assertEqual(page, expected)

    def test_empty_paragraph_and_image_only_note_are_retained(self):
        doc = Document()
        reference(doc.add_paragraph('Texto')._p, 'endnote', 1)
        p = paragraph('Primeiro')
        numbering = element('r')
        numbering.append(element('endnoteRef'))
        numbering.append(element('t', 'NUMERO'))
        p.insert(0, numbering)
        image = element('r')
        image.append(element('pict'))
        image_p = element('p')
        image_p.append(image)
        notes(doc, 'endnote', [(1, [p, element('p'), paragraph('Terceiro')], None),
                               (2, [image_p], None)])
        _, page = self.extract(doc)
        self.assertEqual(page.blocks[1].text, '[^e1]: Primeiro\n\n    \n\n    Terceiro')
        self.assertEqual(page.blocks[2].text, '[^e2]: ')
        self.assertIn('docx_note_complex_content_flattened', page.warnings)

    def test_registry_lists_new_losses_without_aggregate(self):
        losses = registry.lookup('.docx').known_losses
        self.assertNotIn('docx_headers_footnotes_comments_not_extracted', losses)
        self.assertTrue({'docx_note_definition_unreferenced', 'docx_note_reference_unresolved',
                         'docx_note_complex_content_flattened', 'docx_headers_footers_not_extracted',
                         'docx_comments_not_extracted', 'docx_formatting_partial'} <= set(losses))


if __name__ == '__main__':
    unittest.main()
